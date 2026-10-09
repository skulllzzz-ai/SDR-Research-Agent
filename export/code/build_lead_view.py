"""Build the lead's blind scoring workbook: data-exports/lead/out/lead-YYYY-MM-DD-out.xlsx.

    python build_lead_view.py --draw data-exports/master-log/draw-2026-10-01.json \
        --agent-runs data-exports/agent-runs/batch-v1

Whole companies only (DECISIONS.md 12): a company enters the lead's file when both its rows are
finished, and the two arrive in the same release:
  - its Agent alone row, the validated agent row (pass only the v1 batch: one v1 run per company
    serves the agent-alone arm, also after a v2 exists),
  - its SDR alone or Pair row, marked Complete in the latest SDR return.
A company seen once would give away which arm made its row, so a company with one finished row
waits. The arms in a release are whatever its whole companies bring.
The lead sees values only: a blind row code, the company, its market, the research fields, the
verdict. No arm, no work order, no company ID, no colours, no comments, no agent flag, no CRM
status, no notes, no difficulty label. "Not Sure" typed by the SDR in a Y/N product cell is shown
blank (DECISIONS.md 13; it counts as not found either way) and the sheet says a blank cell means
not found. A Dead end row shows only the short-row fields (company, market, verdict, type, reason
with source), whichever arm made it. Rows are shuffled so that a company's two rows never sit next
to each other. The scoring questions and the existing-account rule are printed above the headers.

Each release is built on the latest return in lead/in/ and adds only companies not released before.
The key (row code -> company, arm) lives in master-log/lead-key.json and is never sent.
"""

import argparse
import hashlib
from collections import Counter
import json
import random
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import build_sdr_release as sdr
from contract import (
    ARM_AGENT, ARM_SDR, DIFFICULTIES, FIELD_KEYS, PRODUCT_CATEGORIES, RESEARCH_FIELDS, ROW_AGENT_ALONE, ROW_PAIR,
    ROW_SDR_ALONE, SDR_UNSURE_VALUES, SHEET_RESEARCH, VERDICT_DEAD_END,
)
from xlsx_util import BASE, BOLD, BOX, HEADER_FILL, TITLE, WARN_FILL, WRAP_TOP, gate, seal

HERE = Path(__file__).resolve().parent
SHEET = "SCORING"
SHORT_ROW = ["Verdict", "Dead-end type", "Dead-end reason with source"]
SCORING = ["Acceptable", "Correctness", "Completeness", "Scoring start", "Scoring end", "Comment"]
HEADERS = ["Row code", "Company name", "Market"] + RESEARCH_FIELDS + SHORT_ROW + SCORING
LISTS = {"Acceptable": ["Yes", "No"], "Correctness": ["None", "One", "More than one"], "Completeness": ["Full", "Partial", "Thin"]}
PRINTED = [
    "HOW TO SCORE. For each row, top to bottom:",
    "1. Click Scoring start and press Ctrl+Shift+; (or type the time, e.g. 14:05), then read the row.",
    "2. Fill Acceptable, Correctness and Completeness.",
    "3. Click Scoring end and press Ctrl+Shift+;",
    "Acceptable: Yes or No. Would you let the SDR call from this row?",
    "Correctness: None, One, or More than one wrong fact in the key fields.",
    "Completeness: Full, Partial or Thin.",
    "Use Comment only to explain a score in one line: which field is wrong and why, what is missing, or why the SDR should not call. "
    "Leave it blank otherwise.",
    "A blank cell means not found.",
    "A Dead end row shows only the company, the market, the verdict, the dead-end type and the reason with its source. "
    "Judge it on the verdict and the reason, not on completeness.",
    "Existing accounts: judge every row's verdict on fit as a buyer. A row is not wrong for calling an existing account a prospect, "
    "and a row calling a company an existing active account is right if the CRM says so.",
    "Each company appears twice, on separate rows. Score each row on its own. Please do not add, delete, sort or move rows, "
    "and do not edit the row code, company or the research cells.",
]
HEADER_ROW = len(PRINTED) + 3
# The dead-end re-check (Mayank, 9 Oct; DECISIONS.md 18): the scoring sheet's "would you let the SDR call from this row?"
# was answered literally on dead-end rows, so the lead judges each dead end again on one question, the verdict's own.
RECHECK_SHEET = "RE-CHECK"
RECHECK_QUESTION = "Is this dead end right, on its reason and source? Yes or No"
RECHECK_HEADERS = ["Row code", "Company name", "Market"] + SHORT_ROW + [RECHECK_QUESTION, "Comment"]
RECHECK_PRINTED = [
    "RE-CHECK OF THE DEAD-END ROWS. For each row, answer one question:",
    RECHECK_QUESTION + ".",
    "Judge only the verdict, the dead-end type and the reason with its source: is this company rightly a dead end for us?",
    "Use Comment only to say in one line why, if you answer No. Leave it blank otherwise.",
    "Please do not add, delete, sort or move rows, and do not edit the row code, company or the other cells.",
]
RECHECK_HEADER_ROW = len(RECHECK_PRINTED) + 3
# Wording that would reveal an arm or leak a hidden column. The arm labels are checked as whole
# cells ("pair" alone is also a trade word, as in matched pairs); the rest anywhere in a cell.
ARM_LABELS = [ARM_SDR, ARM_AGENT, ROW_AGENT_ALONE, ROW_PAIR]
FORBIDDEN = [ARM_AGENT, ROW_AGENT_ALONE, "Not attempted", "Agent flag", "CRM status", "Work order"]
UNSURE = {value.lower() for value in SDR_UNSURE_VALUES}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def exchange_files(folder, kind):
    name = re.compile(rf"^lead-(\d{{4}}-\d{{2}}-\d{{2}})-{kind}(?:-(\d+))?\.xlsx$")
    found = [(m.group(1), int(m.group(2) or 1), p) for p in Path(folder).glob("lead-*.xlsx") if (m := name.match(p.name))]
    return [p for _, _, p in sorted(found)]


def clean(value):
    return "" if value is None else str(value).strip()


def agent_values(row):
    values = {name: row["fields"][key]["value"].strip() for name, key in FIELD_KEYS.items()}
    verdict = row["verdict"]
    dead_end = verdict["value"] == VERDICT_DEAD_END
    reason = " ".join(part for part in (verdict["reason"].strip(), verdict["source"].strip()) if part)
    values.update({"Verdict": verdict["value"], "Dead-end type": verdict["dead_end_type"] if dead_end else "",
                   "Dead-end reason with source": reason if dead_end else ""})
    return values


def collect(draw, run_folders, sdr_return):
    """Every finished row: (company, row type, values, source, unsure cells shown blank). An unsure
    value typed in a Y/N product cell is blanked here (DECISIONS.md 13); it is counted only where
    the lead would have seen it (not on a Dead end row, which shows the short-row fields only)."""
    finished, skipped = [], []
    for c in draw["companies"]:
        row = sdr.find_agent_row(c, run_folders)
        if row is not None:
            finished.append((c, ROW_AGENT_ALONE, agent_values(row), "agent run", 0))
    if sdr_return:
        ws = load_workbook(sdr_return)[SHEET_RESEARCH]
        columns = sdr.header_columns(ws)
        by_id = {c["company_id"]: c for c in draw["companies"]}
        for r, cid in sdr.data_rows(ws, columns):
            values = {name: clean(ws.cell(row=r, column=columns[name]).value) for name in RESEARCH_FIELDS + SHORT_ROW}
            if clean(ws.cell(row=r, column=columns["Status"]).value) != "Complete":
                if values["Verdict"]:
                    skipped.append(cid)
                continue
            blanked = 0
            for name in PRODUCT_CATEGORIES:
                if values[name].lower() in UNSURE:
                    values[name] = ""
                    blanked += values["Verdict"] != VERDICT_DEAD_END
            c = by_id[cid]
            finished.append((c, ROW_SDR_ALONE if c["arm"] == ARM_SDR else ROW_PAIR, values, sdr_return.name, blanked))
    return finished, skipped


def is_whole(row_types):
    """A company goes to the lead whole: its Agent alone row and its SDR alone or Pair row."""
    return len(row_types) == 2 and ROW_AGENT_ALONE in row_types and (ROW_SDR_ALONE in row_types or ROW_PAIR in row_types)


def shown(values):
    """What the lead sees of a row: a Dead end keeps the short-row fields only."""
    if values["Verdict"] == VERDICT_DEAD_END:
        return {name: (values[name] if name in SHORT_ROW else "") for name in RESEARCH_FIELDS + SHORT_ROW}
    return dict(values)


def shuffle_apart(rows, previous_company, rng):
    """Shuffle so that no two neighbours (and not the row above the first) share a company."""
    for _ in range(100_000):
        rng.shuffle(rows)
        companies = [previous_company] + [row[0]["company_id"] for row in rows]
        if all(a != b for a, b in zip(companies, companies[1:])):
            return rows
    raise SystemExit("could not keep every company's two rows apart")


def new_sheet(wb):
    ws = wb.active
    ws.title = SHEET
    ws.cell(row=1, column=1, value="Lead scoring sheet").font = TITLE
    for i, line in enumerate(PRINTED, start=2):
        cell = ws.cell(row=i, column=1, value=line)
        cell.font, cell.fill, cell.alignment = (BOLD if i == 2 else BASE), WARN_FILL, WRAP_TOP
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=12)
        ws.row_dimensions[i].height = 28
    widths = {"Row code": 9, "Company name": 32, "Market": 8, "Detailed research": 60, "Parent company": 22,
              "Turnover / revenue / market share": 24, "Contact names": 26, "Emails": 28, "Phones": 20,
              "Dead-end reason with source": 44, "Dead-end type": 20, "Comment": 40}
    for col, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=name)
        cell.font, cell.fill, cell.border, cell.alignment = BOLD, HEADER_FILL, BOX, WRAP_TOP
        ws.column_dimensions[get_column_letter(col)].width = widths.get(name, 13)
    ws.row_dimensions[HEADER_ROW].height = 32
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=3)
    for name, values in LISTS.items():
        letter = get_column_letter(HEADERS.index(name) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        dv.error, dv.errorTitle = f"Pick one of: {', '.join(values)}", name
        ws.add_data_validation(dv)
        dv.add(f"{letter}{HEADER_ROW + 1}:{letter}{HEADER_ROW + 400}")
    return ws


def header_row(ws):
    """The row holding the scoring headers. Found, not assumed: the printed lines above it grew on 8 Oct, so files
    built before then have their headers higher up."""
    for r in range(1, 61):
        if [ws.cell(row=r, column=c).value for c in range(1, len(HEADERS) + 1)] == HEADERS:
            return r
    raise SystemExit("no scoring header row found in the first 60 rows")


def scoring_rows(ws):
    """(row number, row code) of every scoring row in the sheet."""
    return [(r, str(ws.cell(row=r, column=1).value)) for r in range(header_row(ws) + 1, ws.max_row + 1)
            if ws.cell(row=r, column=1).value not in (None, "")]


def logged_releases(data):
    log = Path(data) / "master-log" / "releases.jsonl"
    return [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line.strip()] if log.exists() else []


def lead_releases(data):
    """The lead releases in force, oldest first: a file issued again (--reissue) counts once, as its newest version."""
    superseded = {entry["reissue_of"] for entry in logged_releases(data) if entry.get("reissue_of")}
    return [p for p in exchange_files(Path(data) / "lead" / "out", "out") if p.name not in superseded]


def reissue(data, draw, name, reason):
    """Issues the latest lead release again, unworked, with the printed lines of this version: the same rows, codes,
    order and values under a new name (-2, -3). The first file stays on disk; the release log marks it superseded."""
    releases, returns = lead_releases(data), exchange_files(Path(data) / "lead" / "in", "in")
    if not releases or releases[-1].name != name or len(releases) <= len(returns):
        raise SystemExit(f"REFUSED: only the latest lead release that has not come back is issued again, not {name}")
    old_path = Path(data) / "lead" / "out" / name
    old = load_workbook(old_path)[SHEET]
    rows = scoring_rows(old)
    column = {heading: i + 1 for i, heading in enumerate(HEADERS)}
    scored = sum(1 for r, _ in rows if any(clean(old.cell(row=r, column=column[n]).value) for n in SCORING))
    if scored:
        raise SystemExit(f"REFUSED: {name} carries scores on {scored} rows; a scored file is never issued again")
    m = re.match(r"^lead-(\d{4}-\d{2}-\d{2})-out(?:-(\d+))?\.xlsx$", name)
    out = old_path.with_name(f"lead-{m.group(1)}-out-{int(m.group(2) or 1) + 1}.xlsx")
    if out.exists():
        raise SystemExit(f"{out.name} exists. A release is never overwritten.")
    key = json.loads((Path(data) / "master-log" / "lead-key.json").read_text(encoding="utf-8"))
    by_code = {entry["row_code"]: entry for entry in key["rows"]}
    wb = Workbook()
    ws = new_sheet(wb)
    added = []
    for i, (r, code) in enumerate(rows):
        target = HEADER_ROW + 1 + i
        for col in column.values():
            cell = ws.cell(row=target, column=col, value=old.cell(row=r, column=col).value)
            cell.font, cell.alignment = BASE, Alignment(wrap_text=True, vertical="top")
        for heading in ("Scoring start", "Scoring end"):
            ws.cell(row=target, column=column[heading]).number_format = "hh:mm"
        added.append({**by_code[code], "shown": {n: clean(old.cell(row=r, column=column[n]).value) for n in RESEARCH_FIELDS + SHORT_ROW}})
    wb.save(out)
    print(f"LEAD RELEASE {out.name}: {name} issued again with this version's printed lines; {len(rows)} rows")
    passed = verify(out, key, added, draw)
    same = [code for _, code in scoring_rows(load_workbook(out)[SHEET])] == [code for _, code in rows]
    print(f'{"PASS" if same else "FAIL"}  the rows keep their codes and their order')
    gate_ok = gate(out)
    if not (passed and same and gate_ok):
        failed = out.with_name(out.stem + ".FAILED.xlsx")
        out.rename(failed)
        raise SystemExit(f"the re-issue failed its checks; kept for inspection as {failed.name}, not to be sent")
    entry = {"file": out.name, "sha256": sha256(out), "built_at": datetime.now().astimezone().isoformat(timespec="seconds"),
             "to": "Lead", "built_on": f"{name}, issued again", "rows": len(rows), "new_rows": 0, "reissue_of": name, "reason": reason}
    with open(Path(data) / "master-log" / "releases.jsonl", "a", encoding="utf-8") as handle:
        handle.write(json.dumps(entry) + "\n")
    print(f"\nLOGGED in master-log/releases.jsonl: {entry}")
    print(f'{"SEALED" if seal(out) else "NOT SEALED"}  {out.name} is read-only on disk')


def recheck(data, draw, out):
    """The dead-end rows of the lead's latest return, as the lead saw them, on a sheet with one question: is this dead end
    right, on its reason and source? Same row codes; no scores, no arm, nothing the lead did not see before."""
    releases, returns = lead_releases(data), exchange_files(Path(data) / "lead" / "in", "in")
    if len(releases) > len(returns):
        raise SystemExit(f"{releases[-1].name} has not come back yet. File its return in lead/in/ first.")
    if out.exists():
        raise SystemExit(f"{out.name} exists. A release is never overwritten: pass --seq for another one today.")
    key = {e["row_code"]: e for e in json.loads((Path(data) / "master-log" / "lead-key.json").read_text(encoding="utf-8"))["rows"]}
    source = load_workbook(returns[-1])[SHEET]
    h = header_row(source)
    column = {source.cell(row=h, column=c).value: c for c in range(1, source.max_column + 1)}
    view = ["Row code", "Company name", "Market"] + SHORT_ROW
    rows = [{name: clean(source.cell(row=r, column=column[name]).value) for name in view} for r, _ in scoring_rows(source)]
    dead = [row for row in rows if row["Verdict"] == VERDICT_DEAD_END]
    if not dead:
        raise SystemExit(f"{returns[-1].name} holds no dead-end row: nothing to re-check")

    wb = Workbook()
    ws = wb.active
    ws.title = RECHECK_SHEET
    ws.cell(row=1, column=1, value="Lead re-check sheet").font = TITLE
    for i, line in enumerate(RECHECK_PRINTED, start=2):
        cell = ws.cell(row=i, column=1, value=line)
        cell.font, cell.fill, cell.alignment = (BOLD if i <= 3 else BASE), WARN_FILL, WRAP_TOP
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        ws.row_dimensions[i].height = 22
    widths = [9, 32, 8, 12, 22, 60, 28, 44]
    for col, name in enumerate(RECHECK_HEADERS, start=1):
        cell = ws.cell(row=RECHECK_HEADER_ROW, column=col, value=name)
        cell.font, cell.fill, cell.border, cell.alignment = BOLD, HEADER_FILL, BOX, WRAP_TOP
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.row_dimensions[RECHECK_HEADER_ROW].height = 44
    ws.freeze_panes = ws.cell(row=RECHECK_HEADER_ROW + 1, column=3)
    for i, row in enumerate(dead, start=RECHECK_HEADER_ROW + 1):
        for col, name in enumerate(view, start=1):
            ws.cell(row=i, column=col, value=row[name] or None)
        for col in range(1, len(RECHECK_HEADERS) + 1):
            cell = ws.cell(row=i, column=col)
            cell.font, cell.alignment = BASE, Alignment(wrap_text=True, vertical="top")
    letter = get_column_letter(RECHECK_HEADERS.index(RECHECK_QUESTION) + 1)
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv.error, dv.errorTitle = "Pick Yes or No", "Dead end right?"
    ws.add_data_validation(dv)
    dv.add(f"{letter}{RECHECK_HEADER_ROW + 1}:{letter}{RECHECK_HEADER_ROW + 100}")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"LEAD RELEASE {out.name}: the {len(dead)} dead-end rows of {returns[-1].name}, to be judged on the verdict")

    results = []

    def check(name, passed, detail=""):
        results.append(bool(passed))
        print(f'{"PASS" if passed else "FAIL"}  {name}{"  [" + str(detail) + "]" if detail != "" else ""}')

    wb = load_workbook(out)
    ws = wb[RECHECK_SHEET]
    check("one sheet, the re-check layout, the question in Mayank's words",
          wb.sheetnames == [RECHECK_SHEET] and [ws.cell(row=RECHECK_HEADER_ROW, column=c).value for c in range(1, len(RECHECK_HEADERS) + 1)]
          == RECHECK_HEADERS)
    got = [{name: clean(ws.cell(row=r, column=c).value) for c, name in enumerate(view, start=1)}
           for r in range(RECHECK_HEADER_ROW + 1, ws.max_row + 1) if clean(ws.cell(row=r, column=1).value)]
    check("the dead-end rows of the lead's return, each exactly as the lead saw it, in the same order", got == dead, f"{len(got)} rows")
    check("every row code is in the key; every dead-end row of the return is here, whichever arm made it",
          all(row["Row code"] in key for row in got) and len(got) == len(dead),
          ", ".join(f"{k} {v}" for k, v in sorted(Counter(key[row["Row code"]]["row_type"] for row in got).items())))
    answers = [ws.cell(row=r, column=c).value for r in range(RECHECK_HEADER_ROW + 1, ws.max_row + 1) for c in (7, 8)]
    check("the answer and Comment cells arrive blank", all(v in (None, "") for v in answers))
    cells = [cell for row in ws.iter_rows() for cell in row]
    texts = [str(cell.value) for cell in cells if cell.value is not None]
    hits = {word: sum(word.lower() in text.lower() for text in texts) for word in FORBIDDEN}
    hits.update({f"cell = {word}": sum(text.strip().lower() == word.lower() for text in texts) for word in ARM_LABELS})
    marks = sum(sdr.marker_colour(cell) in sdr.MARKER_COLOURS.values() for cell in cells) + sum(cell.comment is not None for cell in cells)
    check("no arm, flag, CRM or order wording; no marker colour or cell comment", not any(hits.values()) and marks == 0,
          {k: v for k, v in hits.items() if v} or 0)
    names = {c["company_name"] for c in draw["companies"]}
    check("only drawn companies; no company ID or difficulty label",
          all(row["Company name"] in names for row in got)
          and not any(re.fullmatch(r"C\d{2}(-[ASP]\d?)?", t) or t.lower() in DIFFICULTIES for t in texts))
    fonts = {cell.font.name for cell in cells if cell.value is not None}
    check("Arial only, no formulas", fonts == {"Arial"} and not any(t.startswith("=") for t in texts), sorted(fonts))
    gate_ok = gate(out)
    print(f"\n{sum(results)} of {len(results)} checks passed")
    if not (all(results) and gate_ok):
        failed = out.with_name(out.stem + ".FAILED.xlsx")
        out.rename(failed)
        raise SystemExit(f"the re-check sheet failed its checks; kept for inspection as {failed.name}, not to be sent")
    entry = {"file": out.name, "sha256": sha256(out), "built_at": datetime.now().astimezone().isoformat(timespec="seconds"),
             "to": "Lead", "kind": "dead-end re-check", "built_on": returns[-1].name, "rows": len(dead), "new_rows": 0,
             "rows_by_arm": dict(Counter(key[row["Row code"]]["row_type"] for row in dead)), "question": RECHECK_QUESTION}
    with open(Path(data) / "master-log" / "releases.jsonl", "a", encoding="utf-8") as handle:
        handle.write(json.dumps(entry) + "\n")
    print(f"\nLOGGED in master-log/releases.jsonl: {entry}")
    print(f'{"SEALED" if seal(out) else "NOT SEALED"}  {out.name} is read-only on disk')


def verify(out, key, added, draw):
    wb = load_workbook(out)
    ws = wb[SHEET]
    results = []

    def check(name, passed, detail=""):
        results.append(bool(passed))
        print(f'{"PASS" if passed else "FAIL"}  {name}{"  [" + str(detail) + "]" if detail != "" else ""}')

    headers = [ws.cell(row=HEADER_ROW, column=c).value for c in range(1, len(HEADERS) + 1)]
    column = {name: i + 1 for i, name in enumerate(headers)}
    check("headers are the scoring layout", headers == HEADERS and wb.sheetnames == [SHEET], f"{len(headers)} columns")
    rows = scoring_rows(ws)
    by_code = {entry["row_code"]: entry for entry in key["rows"]}
    check("every row in the file has a key entry, and every key entry a row", sorted(code for _, code in rows) == sorted(by_code),
          f"{len(rows)} rows")
    companies = [by_code[code]["company_id"] for _, code in rows if code in by_code]
    check("no company's two rows sit next to each other", all(a != b for a, b in zip(companies, companies[1:])))
    per_company = {}
    for _, code in rows:
        if code in by_code:
            per_company.setdefault(by_code[code]["company_id"], []).append(by_code[code]["row_type"])
    check("whole companies only: each has its agent row and its SDR or pair row, so no company appears once",
          all(is_whole(kinds) for kinds in per_company.values()), f"{len(per_company)} companies")
    yes_no = [clean(ws.cell(row=r, column=column[name]).value) for r, _ in rows for name in PRODUCT_CATEGORIES]
    check("Y/N product cells hold Y, N or blank only: no 'Not Sure' shown", all(v in ("", "Y", "N") for v in yes_no),
          f"{sum(v not in ('', 'Y', 'N') for v in yes_no)} other values")

    cells = [cell for row in ws.iter_rows() for cell in row]
    marker_cells = sum(sdr.marker_colour(cell) in sdr.MARKER_COLOURS.values() for cell in cells)
    comments = sum(cell.comment is not None for cell in cells)
    check("planted-marker control: marker colours in the file", marker_cells == 0, marker_cells)
    check("planted-marker control: comments (sources) in the file", comments == 0, comments)
    texts = [str(cell.value) for cell in cells if cell.value is not None]
    hits = {word: sum(word.lower() in text.lower() for text in texts) for word in FORBIDDEN}
    hits.update({f"cell = {word}": sum(text.strip().lower() == word.lower() for text in texts) for word in ARM_LABELS})
    check("no arm, flag, CRM or order wording anywhere", not any(hits.values()), {k: v for k, v in hits.items() if v} or 0)
    data_texts = [clean(ws.cell(row=r, column=c).value) for r, _ in rows for c in range(1, len(HEADERS) + 1)]
    ids = sum(bool(re.fullmatch(r"C\d{2}(-[ASP]\d?)?", text)) for text in data_texts)
    labels = sum(text.lower() in DIFFICULTIES for text in data_texts)
    check("no company ID and no difficulty label in any row", ids == 0 and labels == 0, f"{ids} ids, {labels} labels")

    dead = [r for r, _ in rows if clean(ws.cell(row=r, column=column["Verdict"]).value) == VERDICT_DEAD_END]
    filled = sum(bool(clean(ws.cell(row=r, column=column[name]).value)) for r in dead for name in RESEARCH_FIELDS)
    short = all(clean(ws.cell(row=r, column=column[name]).value) for r in dead for name in ["Company name", "Market"] + SHORT_ROW)
    check("Dead end rows show only the short-row fields", filled == 0 and short, f"{len(dead)} dead-end rows, {filled} extra cells")
    new_codes = {entry["row_code"] for entry in added}
    faithful = 0
    for r, code in rows:
        if code in new_codes:
            want = next(entry for entry in added if entry["row_code"] == code)["shown"]
            faithful += all(clean(ws.cell(row=r, column=column[name]).value) == want[name] for name in RESEARCH_FIELDS + SHORT_ROW)
    check("every new row shows exactly its source values", faithful == len(added), f"{faithful} of {len(added)}")
    blank_scores = all(not clean(ws.cell(row=r, column=column[name]).value) for r, code in rows if code in new_codes for name in SCORING)
    stamps = all(ws.cell(row=r, column=column[name]).number_format == "hh:mm" for r, _ in rows for name in ("Scoring start", "Scoring end"))
    check("each row has its two scoring-time stamp cells, and new rows arrive unscored", blank_scores and stamps)
    printed = " ".join(texts)
    check("the scoring questions and the existing-account rule are printed on the sheet", all(line in printed for line in PRINTED))
    names = {c["company_name"] for c in draw["companies"]}
    outsiders = [r["company_name"] for r in draw["not_drawn"] + draw["pasted_reserves"]]
    check("only drawn companies appear", all(clean(ws.cell(row=r, column=column["Company name"]).value) in names for r, _ in rows)
          and not any(name.lower() in text.lower() for name in outsiders for text in texts))
    fonts = {cell.font.name for cell in cells if cell.value is not None}
    formulas = sum(text.startswith("=") for text in texts)
    check("Arial only, no formulas", fonts == {"Arial"} and formulas == 0, f"{sorted(fonts)}, {formulas} formulas")
    print(f"\n{sum(results)} of {len(results)} checks passed")
    return all(results)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--draw", required=True)
    parser.add_argument("--agent-runs", nargs="*", default=[])
    parser.add_argument("--date", default=datetime.now().strftime("%Y-%m-%d"))
    parser.add_argument("--seq", type=int)
    parser.add_argument("--seed", type=int, default=20261001)
    parser.add_argument("--data-root", default=str(HERE / "data-exports"))
    parser.add_argument("--reissue", help="the latest lead release, unworked, issued again with this version's printed lines")
    parser.add_argument("--reason", help="--reissue: why it is issued again (recorded in the release log)")
    parser.add_argument("--recheck-dead-ends", action="store_true",
                        help="the dead-end rows of the lead's latest return, to be judged on the verdict alone (DECISIONS.md 18)")
    args = parser.parse_args()
    data = Path(args.data_root)
    draw = json.loads(Path(args.draw).read_text(encoding="utf-8"))
    if args.reissue:
        assert args.reason, "--reissue needs --reason"
        reissue(data, draw, args.reissue, args.reason)
        return

    out = data / "lead" / "out" / f'lead-{args.date}-out{f"-{args.seq}" if args.seq else ""}.xlsx'
    if args.recheck_dead_ends:
        recheck(data, draw, out)
        return
    if out.exists():
        raise SystemExit(f"{out.name} exists. A release is never overwritten: pass --seq for another one today.")
    releases, returns = lead_releases(data), exchange_files(data / "lead" / "in", "in")
    if len(releases) > len(returns):
        raise SystemExit(f"{releases[-1].name} has not come back yet. File its return in lead/in/ first.")

    key_path = data / "master-log" / "lead-key.json"
    key = json.loads(key_path.read_text(encoding="utf-8")) if key_path.exists() else {"seed": args.seed, "rows": []}
    sdr_returns = sdr.exchange_files(data / "sdr" / "in", "in")
    finished, skipped = collect(draw, args.agent_runs, sdr_returns[-1] if sdr_returns else None)
    released = Counter(entry["company_id"] for entry in key["rows"])
    assert all(n == 2 for n in released.values()), "the lead key holds a company with one row: companies go to the lead whole"
    by_company = {}
    for row in finished:
        by_company.setdefault(row[0]["company_id"], []).append(row)
    whole = sorted(cid for cid, rows in by_company.items() if cid not in released and is_whole([r[1] for r in rows]))
    waiting = sorted(cid for cid, rows in by_company.items() if cid not in released and cid not in whole)
    fresh = [row for cid in whole for row in by_company[cid]]
    if not fresh:
        raise SystemExit(f"no whole company is waiting to be scored: a company goes to the lead only when both its rows "
                         f"are finished ({len(waiting)} companies have one so far)")

    if returns:
        wb = load_workbook(returns[-1])
        ws = wb[SHEET]
        existing = scoring_rows(ws)
        assert sorted(code for _, code in existing) == sorted(e["row_code"] for e in key["rows"]), \
            "the returned file's rows do not match the key: rows were added, deleted or re-coded"
        next_row = max(r for r, _ in existing) + 1
        previous = next(e["company_id"] for e in key["rows"] if e["row_code"] == existing[-1][1])
    else:
        wb = Workbook()
        ws = new_sheet(wb)
        next_row, previous = HEADER_ROW + 1, None

    rng = random.Random(args.seed + len(key["rows"]))
    fresh = shuffle_apart(fresh, previous, rng)
    used = {entry["row_code"] for entry in key["rows"]}
    codes = [f"R{n}" for n in rng.sample(range(100, 1000), len(used) + len(fresh)) if f"R{n}" not in used][:len(fresh)]
    column = {name: i + 1 for i, name in enumerate(HEADERS)}
    added = []
    for (company, row_type, values, source, blanked), code in zip(fresh, codes):
        visible = shown(values)
        ws.cell(row=next_row, column=column["Row code"], value=code)
        ws.cell(row=next_row, column=column["Company name"], value=company["company_name"])
        ws.cell(row=next_row, column=column["Market"], value=company["market"])
        for name in RESEARCH_FIELDS + SHORT_ROW:
            ws.cell(row=next_row, column=column[name], value=visible[name] or None)
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=next_row, column=col)
            cell.font, cell.alignment = BASE, Alignment(wrap_text=True, vertical="top")
        for name in ("Scoring start", "Scoring end"):
            ws.cell(row=next_row, column=column[name]).number_format = "hh:mm"
        entry = {"row_code": code, "company_id": company["company_id"], "row_type": row_type, "source": source, "released_in": out.name,
                 "unsure_shown_blank": blanked}
        key["rows"].append(entry)
        added.append({**entry, "shown": visible})
        next_row += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    print(f"LEAD RELEASE {out.name}  built on {returns[-1].name if returns else 'a new sheet'}: "
          f"{len(added)} new rows ({len(whole)} whole companies), {len(key['rows'])} in all; "
          f"{len(waiting)} companies wait for their second row")
    if skipped:
        print(f"NOTE  {len(skipped)} SDR rows carry a verdict but are not marked Complete, so they were left out")
    passed = verify(out, key, added, draw)
    gate_ok = gate(out)
    if not (passed and gate_ok):
        failed = out.with_name(out.stem + ".FAILED.xlsx")
        out.rename(failed)
        raise SystemExit(f"lead release failed its checks; kept for inspection as {failed.name}, not to be sent")
    key_path.write_text(json.dumps(key, indent=2) + "\n", encoding="utf-8")
    kinds = {kind: sum(e["row_type"] == kind for e in added) for kind in (ROW_AGENT_ALONE, ROW_SDR_ALONE, ROW_PAIR)}
    entry = {"file": out.name, "sha256": sha256(out), "built_at": datetime.now().astimezone().isoformat(timespec="seconds"),
             "to": "Lead", "built_on": returns[-1].name if returns else "new sheet", "rows": len(key["rows"]), "new_rows": len(added),
             "new_rows_by_arm": kinds, "new_companies": len(whole), "companies_waiting_for_second_row": len(waiting),
             "unsure_cells_shown_blank": sum(e["unsure_shown_blank"] for e in added),
             "agent_runs": [Path(folder).name for folder in args.agent_runs]}
    with open(data / "master-log" / "releases.jsonl", "a", encoding="utf-8") as handle:
        handle.write(json.dumps(entry) + "\n")
    print(f"\nLOGGED in master-log/releases.jsonl: {entry}")
    print(f'{"SEALED" if seal(out) else "NOT SEALED"}  {out.name} is read-only on disk')


if __name__ == "__main__":
    main()
