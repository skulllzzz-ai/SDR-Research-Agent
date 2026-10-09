"""Build the next SDR workbook release: data-exports/sdr/out/sdr-YYYY-MM-DD-out.xlsx.

    python build_sdr_release.py --draw data-exports/master-log/draw-2026-10-01.json --through-order 7
    python build_sdr_release.py --draw ... --through-order 12 --agent-runs data-exports/agent-runs/<batch>

The release holds every company whose work order is at or below --through-order. It is built on
the latest return in sdr/in/ (sdr-YYYY-MM-DD-in.xlsx), so nobody merges files; with nothing sent
yet it is built on the blank template, with the EXAMPLE row removed. Rows already in the base are
never touched. A release is never overwritten, and a new one is refused while the last one has not
come back: a release has come back when a return filed with sdr_return.py answers it
(master-log/returns.jsonl). A correction of a return is a second file for the same release, so the
files in sdr/in/ are not counted.

A With-agent company is added only once its validated agent row exists (--agent-runs): the row
arrives filled, with each field coloured by the agent's marker (green sure, amber guessed, red not
found), the source as a cell comment, and the flag in the dark grey band. The verdict and CRM
cells stay empty for the SDR. The release log names, for each agent row added, its run folder, its
hash, the prompt version, model and effort, because the pair arm is reported by agent version.

After saving, the file is re-read from disk and checked: rows held, nothing from an unreleased
company, no difficulty label, and the contamination count (agent content on SDR-alone companies,
which must be 0). Then the spreadsheet gate runs. A file that fails is renamed .FAILED and is not
to be sent.
"""

import argparse
import hashlib
import json
import re
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

from contract import (
    AGENT_FLAG, ARM_AGENT, ARM_SDR, DIFFICULTIES, FIELD_KEYS, FIRST_DATA_ROW, HEADER_ROW, LAST_DATA_ROW,
    RESEARCH_COLUMNS, RESEARCH_FIELDS, SET_BY_MAYANK, SHEET_INSTRUCTIONS, SHEET_LISTS,
    SHEET_RESEARCH, SHEET_SNAPSHOT, VERDICT_COLUMNS,
)
from validate_row import validate
from xlsx_util import GUESSED_FILL, NOT_FOUND_FILL, SURE_FILL, gate, seal

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "templates" / "sdr-workbook-template-v3.xlsx"
MARKER_FILL = {"sure": SURE_FILL, "guessed": GUESSED_FILL, "not found": NOT_FOUND_FILL}
MARKER_COLOURS = {marker: fill.start_color.rgb[-6:] for marker, fill in MARKER_FILL.items()}

# Two instruction lines in the blank template go stale once the example row is removed and the
# calendar has moved. Left as they were, the first tells the SDR to delete a row, which would now
# be a real company. Exact old text -> new text.
INSTRUCTION_FIXES = {
    "The grey italic EXAMPLE row on the RESEARCH sheet shows the format. Delete it before the first real company.":
        "Your companies are already listed on the RESEARCH sheet, from row 3 down. The example row has been removed, so do not delete any row.",
    "STEPS FOR A COMPANY MARKED WITH AGENT (from Thursday)":
        "STEPS FOR A COMPANY MARKED WITH AGENT (these companies arrive in a later file)",
}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def header_columns(ws):
    """Column number by header text, read from the sheet every time (never by remembered letter)."""
    return {ws.cell(row=HEADER_ROW, column=c).value: c for c in range(1, ws.max_column + 1)
            if ws.cell(row=HEADER_ROW, column=c).value}


def exchange_files(folder, kind):
    """The releases ('out') or returns ('in') in a folder, oldest first."""
    name = re.compile(rf"^sdr-(\d{{4}}-\d{{2}}-\d{{2}})-{kind}(?:-(\d+))?\.xlsx$")
    found = [(m.group(1), int(m.group(2) or 1), p) for p in Path(folder).glob("sdr-*.xlsx") if (m := name.match(p.name))]
    return [p for _, _, p in sorted(found)]


def data_rows(ws, columns):
    """(row number, company id) for every data row that holds a company."""
    found = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        company_id = ws.cell(row=r, column=columns["Company ID"]).value
        if company_id not in (None, ""):
            found.append((r, str(company_id)))
    return found


def clear_example_row(ws, columns):
    name = ws.cell(row=FIRST_DATA_ROW, column=columns["Company name"]).value
    assert name and "EXAMPLE" in str(name), "row 3 of the template is not the EXAMPLE row"
    for col in range(1, len(RESEARCH_COLUMNS) + 1):
        cell = ws.cell(row=FIRST_DATA_ROW, column=col)
        model = ws.cell(row=FIRST_DATA_ROW + 1, column=col)
        cell.value = None
        cell.comment = None
        cell.font, cell.fill, cell.border = copy(model.font), copy(model.fill), copy(model.border)
        cell.alignment, cell.number_format = copy(model.alignment), model.number_format
    ws.row_dimensions[FIRST_DATA_ROW].height = ws.row_dimensions[FIRST_DATA_ROW + 1].height


# Rules added on 5 Oct 2026 after the first return (DECISIONS.md, decisions 1 to 5). Printed once at
# the end of the Instructions sheet of every release from then on, so the SDR reads them where the
# work happens. Exact text: the self-test and the release checks look for these lines.
RULES_ADDED = [
    None,
    "RULES ADDED ON 5 OCT 2026, AFTER THE FIRST RETURN. Please read before the next company.",
    "1. Time stamps are per company and per step. Stamp Step 1 start and Step 1 end when you begin and finish the first "
    "pass on THAT company, and Step 2 start and Step 2 end when you begin and finish the second pass. Never copy the "
    "Step 1 times into Step 2, and never use one start and one end for several companies.",
    "2. After Step 1 on an SDR-alone company, copy that whole row as values to the sheet 'Pass 1 snapshot (paste only)', "
    "then carry on with Step 2 on the RESEARCH sheet.",
    "3. The Y/N product cells (Gold, Natural diamond, LGD, Gemstone, Pearl, Watch, Jewellery) take Y or N only. If you "
    "cannot tell, leave the cell blank. Do not type 'Not Sure'.",
    "4. A Dead end needs one of the four kinds from the dropdown: Competitor, Wrong target account, Wrong segment, "
    "Existing active account. 'No information found' is not a dead end: the company stays a Prospect and the cells "
    "you could not fill stay blank.",
    "5. The grey columns (Order, Company ID, Company name, Market, Arm) are set for you. Do not edit or clear them.",
]


def add_rules(ws):
    """Appends RULES_ADDED below the last instruction line, once. Returns the rows written."""
    present = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    if RULES_ADDED[1] in present:
        return []
    last = max(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value not in (None, ""))
    model = ws.cell(row=last, column=1)
    rows = []
    for i, line in enumerate(RULES_ADDED):
        r = last + 1 + i
        cell = ws.cell(row=r, column=1, value=line)
        cell.font, cell.alignment = copy(model.font), copy(model.alignment)
        if i == 1:
            cell.font = copy(ws.cell(row=1, column=1).font)
        rows.append(r)
    return rows


def fix_instructions(ws):
    fixed = 0
    for r in range(1, ws.max_row + 1):
        cell = ws.cell(row=r, column=1)
        if cell.value in INSTRUCTION_FIXES:
            cell.value = INSTRUCTION_FIXES[cell.value]
            fixed += 1
    assert fixed == len(INSTRUCTION_FIXES), "an instruction line to fix was not found in the template"


def agent_row_path(company, run_folders):
    """<run folder>/<company id>/row.json from the last folder that has one, or None."""
    for folder in reversed(run_folders):
        path = Path(folder) / company["company_id"] / "row.json"
        if path.exists():
            return path
    return None


def find_agent_row(company, run_folders):
    """The validated agent row for a company: <run folder>/<company id>/row.json, last folder wins."""
    path = agent_row_path(company, run_folders)
    if path is None:
        return None
    row = json.loads(path.read_text(encoding="utf-8"))
    problems = validate(row, {k: company[k] for k in ("company_id", "company_name", "market")})
    assert not problems, f"{path} does not pass the validator: {problems}"
    return row


def agent_row_record(company, run_folders):
    """Where a released agent row came from, for the release log: the pair arm is reported by agent
    version (v1, v2), so each row names its run folder, hash, prompt version, model and effort."""
    path = agent_row_path(company, run_folders)
    run_path = path.parent / "run.json"
    run = json.loads(run_path.read_text(encoding="utf-8")) if run_path.exists() else {}
    return {"company_id": company["company_id"], "row": f"{path.parent.parent.name}/{path.parent.name}/row.json",
            "row_sha256": sha256(path), "prompt_version": run.get("prompt_version"), "model": run.get("model_asked"),
            "effort": run.get("effort"), "agent_minutes": run.get("agent_minutes"), "retry_count": run.get("retry_count")}


def filed_returns(data):
    """The filings in master-log/returns.jsonl, oldest first (written by sdr_return.py)."""
    log = Path(data) / "master-log" / "returns.jsonl"
    return [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line.strip()] if log.exists() else []


def write_agent_row(ws, columns, r, row):
    flag = row["flag"]
    ws.cell(row=r, column=columns["Agent flag"], value="Yes" if flag["likely_dead_end"] == "yes" else "No")
    ws.cell(row=r, column=columns["Agent flag reason"], value=flag["reason"].strip() or None)
    ws.cell(row=r, column=columns["Agent flag source"], value=flag["source"].strip() or None)
    for name, key in FIELD_KEYS.items():
        entry, cell = row["fields"][key], ws.cell(row=r, column=columns[name])
        cell.value = entry["value"].strip() or None
        cell.fill = copy(MARKER_FILL[entry["marker"]])
        if entry["source"].strip():
            label = "Source" if entry["marker"] == "sure" else "The guess rests on"
            cell.comment = Comment(f'{label}: {entry["source"].strip()}', "Agent")
            cell.comment.width, cell.comment.height = 320, 90


def marker_colour(cell):
    rgb = cell.fill.start_color.rgb if cell.fill and cell.fill.fill_type == "solid" else None
    return rgb[-6:] if isinstance(rgb, str) else ""


def agent_content_on_sdr_alone(ws, columns):
    """The contamination count: SDR-alone rows carrying anything only an agent row carries
    (a flag cell, a marker colour, a source comment)."""
    dirty = 0
    for r, _ in data_rows(ws, columns):
        if ws.cell(row=r, column=columns["Arm"]).value != ARM_SDR:
            continue
        flag = any(ws.cell(row=r, column=columns[name]).value not in (None, "") for name in AGENT_FLAG)
        marks = any(marker_colour(ws.cell(row=r, column=columns[name])) in MARKER_COLOURS.values()
                    or ws.cell(row=r, column=columns[name]).comment is not None for name in RESEARCH_FIELDS)
        dirty += flag or marks
    return dirty


def dropdown_map(ws):
    """Column -> (first row, last row, list). Excel merges identical dropdowns into one entry when
    it saves, so the comparison is by what each column offers, not by the number of entries."""
    covered = {}
    for dv in ws.data_validations.dataValidation:
        for cell_range in dv.sqref.ranges:
            for col in range(cell_range.min_col, cell_range.max_col + 1):
                covered[col] = (cell_range.min_row, cell_range.max_row, dv.formula1)
    return covered


def all_text(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    yield ws.title, cell.coordinate, cell.value


def font_outliers(wb, base):
    """Cells holding a value in a font other than Arial, split in two: those this release wrote (must
    be none) and those kept exactly as the base had them, value and font (the SDR's own typing or
    pasting in a return; rows already in the base are never touched, formatting included)."""
    ours, kept = [], []
    for ws in wb.worksheets:
        before = base[ws.title] if ws.title in base.sheetnames else None
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or cell.font.name == "Arial":
                    continue
                old = before[cell.coordinate] if before is not None else None
                same = old is not None and old.value == cell.value and old.font.name == cell.font.name
                (kept if same else ours).append((ws.title, cell.coordinate))
    return ours, kept


def sheet_values(ws):
    return {cell.coordinate: cell.value for row in ws.iter_rows() for cell in row if cell.value is not None}


def verify(out_path, base_path, base_is_template, expected, draw, agent_rows, restored=(), rules_rows=()):
    """Re-read the saved file and check it. Returns (all passed, contamination count, the SDR's cells
    kept in a font other than Arial)."""
    wb, base = load_workbook(out_path), load_workbook(base_path)
    ws = wb[SHEET_RESEARCH]
    columns = header_columns(ws)
    results = []

    def check(name, passed, detail=""):
        results.append(bool(passed))
        print(f'{"PASS" if passed else "FAIL"}  {name}{"  [" + str(detail) + "]" if detail != "" else ""}')

    for title in (SHEET_RESEARCH, SHEET_SNAPSHOT):
        headers = [wb[title].cell(row=HEADER_ROW, column=c).value for c in range(1, len(RESEARCH_COLUMNS) + 1)]
        check(f"{title}: the 46 headers match the contract", headers == RESEARCH_COLUMNS)
    check("sheets are the template's four, in order",
          wb.sheetnames == [SHEET_INSTRUCTIONS, SHEET_RESEARCH, SHEET_SNAPSHOT, SHEET_LISTS], wb.sheetnames)
    dropdowns = dropdown_map(ws)
    check("the 15 dropdown columns offer the template's lists on rows 3 to 402",
          dropdowns == dropdown_map(load_workbook(TEMPLATE)[SHEET_RESEARCH]) and len(dropdowns) == 15, f"{len(dropdowns)} columns")
    check("merged cells and frozen panes are unchanged",
          sorted(map(str, ws.merged_cells.ranges)) == sorted(map(str, base[SHEET_RESEARCH].merged_cells.ranges))
          and ws.freeze_panes == base[SHEET_RESEARCH].freeze_panes, ws.freeze_panes)
    formulas = [c for _, c, v in all_text(wb) if v.startswith("=")]
    check("no formulas", not formulas, len(formulas))
    ours, kept = font_outliers(wb, base)
    check("Arial only in every cell this release wrote; cells the return brought in another font are kept as returned",
          not ours, f"{len(ours)} written by this release, {len(kept)} kept from the return")

    rows = dict((cid, r) for r, cid in data_rows(ws, columns))
    held = [tuple(ws.cell(row=r, column=columns[name]).value for name in SET_BY_MAYANK) for r, _ in data_rows(ws, columns)]
    want = [(c["work_order"], c["company_id"], c["company_name"], c["market"], c["arm"]) for c in expected]
    check("rows held = the released companies, in work order", held == want, f"{len(held)} rows")
    arms = [row[4] for row in held]
    print(f"      SDR-alone rows: {arms.count(ARM_SDR)}, With-agent rows: {arms.count(ARM_AGENT)}")

    released = {c["company_name"] for c in expected}
    others = [c["company_name"] for c in draw["companies"] if c["company_name"] not in released]
    others += [r["company_name"] for r in draw["not_drawn"] + draw["pasted_reserves"]]
    texts = list(all_text(wb))
    leaks = [(sheet, at) for sheet, at, value in texts if any(name.lower() in value.lower() for name in others)]
    check("no unreleased, not-drawn or reserve company appears anywhere", not leaks, f"{len(leaks)} hits against {len(others)} names")
    labels = [(sheet, at) for sheet, at, value in texts if value.strip().lower() in DIFFICULTIES]
    check("no difficulty label on any sheet", not labels, len(labels))
    check("the template's EXAMPLE row is gone from the RESEARCH sheet",
          not [at for sheet, at, value in texts if sheet == SHEET_RESEARCH and ("EXAMPLE ROW" in value or value == "EX00")])

    dirty = agent_content_on_sdr_alone(ws, columns)
    check("contamination count: agent content on SDR-alone companies", dirty == 0, dirty)

    new_pairs = [c for c in expected if c.get("_new") and c["arm"] == ARM_AGENT]
    if new_pairs:
        faithful = coloured = commented = empty_verdict = 0
        for c in new_pairs:
            r, row = rows[c["company_id"]], agent_rows[c["company_id"]]
            entries = [(ws.cell(row=r, column=columns[name]), row["fields"][key]) for name, key in FIELD_KEYS.items()]
            faithful += all((cell.value or "") == entry["value"].strip() for cell, entry in entries) and \
                ws.cell(row=r, column=columns["Agent flag"]).value == ("Yes" if row["flag"]["likely_dead_end"] == "yes" else "No")
            coloured += all(marker_colour(cell) == MARKER_COLOURS[entry["marker"]] for cell, entry in entries)
            commented += all((cell.comment is not None) == bool(entry["source"].strip()) for cell, entry in entries) and \
                all(entry["source"].strip() in cell.comment.text for cell, entry in entries if cell.comment)
            empty_verdict += all(ws.cell(row=r, column=columns[name]).value in (None, "") for name in VERDICT_COLUMNS)
        n = len(new_pairs)
        check("new With-agent rows: values and flag equal the validated agent row", faithful == n, f"{faithful} of {n}")
        check("new With-agent rows: every research field is coloured by its marker", coloured == n, f"{coloured} of {n}")
        check("new With-agent rows: a source comment exactly where the agent gave a source", commented == n, f"{commented} of {n}")
        check("new With-agent rows: verdict and CRM cells left empty for the SDR", empty_verdict == n, f"{empty_verdict} of {n}")

    new_rows = {rows[c["company_id"]] for c in expected if c.get("_new")}
    allowed = {(SHEET_RESEARCH, ws.cell(row=r, column=col).coordinate) for r in new_rows for col in range(1, len(RESEARCH_COLUMNS) + 1)}
    allowed |= {(SHEET_RESEARCH, ws.cell(row=r, column=columns[name]).coordinate) for r, name in restored}
    allowed |= {(SHEET_INSTRUCTIONS, f"A{r}") for r in rules_rows}
    instructions = [wb[SHEET_INSTRUCTIONS].cell(row=r, column=1).value for r in range(1, wb[SHEET_INSTRUCTIONS].max_row + 1)]
    check("the rules added on 5 Oct are printed once on the Instructions sheet",
          instructions.count(RULES_ADDED[1]) == 1 and all(line in instructions for line in RULES_ADDED[1:]))
    check("grey cells restored from the draw: only blanks, each now equal to the draw",
          all(ws.cell(row=r, column=columns[name]).value not in (None, "") for r, name in restored), len(restored))
    if base_is_template:
        allowed |= {(SHEET_RESEARCH, ws.cell(row=FIRST_DATA_ROW, column=col).coordinate) for col in range(1, len(RESEARCH_COLUMNS) + 1)}
        allowed |= {(SHEET_INSTRUCTIONS, f"A{r}") for r in range(1, base[SHEET_INSTRUCTIONS].max_row + 1)
                    if base[SHEET_INSTRUCTIONS].cell(row=r, column=1).value in INSTRUCTION_FIXES}
    changed = set()
    for title in wb.sheetnames:
        before, after = sheet_values(base[title]), sheet_values(wb[title])
        changed |= {(title, at) for at in set(before) | set(after) if before.get(at) != after.get(at)}
    check("every changed cell is one this release was meant to change", changed <= allowed,
          f"{len(changed)} changed, {len(changed - allowed)} unexpected")
    for title in (SHEET_SNAPSHOT, SHEET_LISTS):
        check(f"{title}: untouched", not {c for c in changed if c[0] == title})
    if base_is_template:
        def look(row, col):
            font = ws.cell(row=row, column=col).font
            return font.italic, str(font.color.rgb) if font.color else None

        plain = all(look(FIRST_DATA_ROW, col) == look(FIRST_DATA_ROW + 1, col) for col in range(1, len(RESEARCH_COLUMNS) + 1))
        stray = sum(1 for row in ws.iter_rows() for cell in row if cell.comment and cell.row not in new_rows)
        check("row 3 carries no example formatting, and no comment sits outside a new agent row", plain and stray == 0,
              f"{stray} stray comments")
    print(f"\n{sum(results)} of {len(results)} checks passed")
    return all(results), dirty, len(kept)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--draw", required=True)
    parser.add_argument("--through-order", type=int, required=True)
    parser.add_argument("--agent-runs", nargs="*", default=[], help="agent batch folders holding <company id>/row.json")
    parser.add_argument("--date", default=datetime.now().strftime("%Y-%m-%d"))
    parser.add_argument("--seq", type=int, help="second or later release on the same day: -2, -3")
    parser.add_argument("--data-root", default=str(HERE / "data-exports"), help="for tests on made-up data")
    args = parser.parse_args()
    data = Path(args.data_root)

    draw = json.loads(Path(args.draw).read_text(encoding="utf-8"))
    expected = sorted((dict(c) for c in draw["companies"] if c["work_order"] <= args.through_order),
                      key=lambda c: c["work_order"])
    out = data / "sdr" / "out" / f'sdr-{args.date}-out{f"-{args.seq}" if args.seq else ""}.xlsx'
    if out.exists():
        raise SystemExit(f"{out.name} exists. A release is never overwritten: pass --seq for another one today.")
    releases, returns = exchange_files(data / "sdr" / "out", "out"), exchange_files(data / "sdr" / "in", "in")
    if releases and releases[-1].name not in {entry.get("answers") for entry in filed_returns(data)}:
        raise SystemExit(f"{releases[-1].name} has not come back yet. File its return in sdr/in/ with sdr_return.py first: "
                         "each release is built on the latest return.")

    base_is_template = not returns
    base_path = TEMPLATE if base_is_template else returns[-1]
    wb = load_workbook(base_path)
    ws = wb[SHEET_RESEARCH]
    columns = header_columns(ws)
    assert [ws.cell(row=HEADER_ROW, column=c).value for c in range(1, len(RESEARCH_COLUMNS) + 1)] == RESEARCH_COLUMNS, \
        "the base workbook's headers are not the contract's 46 columns"
    if base_is_template:
        clear_example_row(ws, columns)
        fix_instructions(wb[SHEET_INSTRUCTIONS])
    rules_rows = add_rules(wb[SHEET_INSTRUCTIONS])

    present = dict((cid, r) for r, cid in data_rows(ws, columns))
    by_id = {c["company_id"]: c for c in draw["companies"]}
    restored = []
    for cid, r in present.items():
        assert cid in by_id, f"row {r} holds an unknown company id {cid}"
        c = by_id[cid]
        want = dict(zip(SET_BY_MAYANK, (c["work_order"], cid, c["company_name"], c["market"], c["arm"])))
        for name, value in want.items():
            held = ws.cell(row=r, column=columns[name]).value
            if held == value:
                continue
            # a grey cell the SDR cleared by accident is restored from the draw (DECISIONS.md, 5 Oct);
            # a grey cell holding a different value is a real conflict and stops the build
            assert held in (None, ""), f"row {r}: the grey cell {name} holds {held!r} where the draw says {value!r}"
            ws.cell(row=r, column=columns[name], value=value)
            restored.append((r, name))
    if restored:
        print(f"RESTORED {len(restored)} grey cell(s) the return had left blank, from the draw: "
              f"{sorted({name for _, name in restored})}")
    next_row = max(present.values(), default=FIRST_DATA_ROW - 1) + 1
    agent_rows = {}
    for c in expected:
        c["_new"] = c["company_id"] not in present
        if not c["_new"]:
            continue
        if c["arm"] == ARM_AGENT:
            agent_rows[c["company_id"]] = find_agent_row(c, args.agent_runs)
            if agent_rows[c["company_id"]] is None:
                raise SystemExit(f'{c["company_id"]} is a With-agent company with no validated agent row in {args.agent_runs}. '
                                 "Its row is released only once its agent row exists.")
        assert next_row <= LAST_DATA_ROW
        for name, value in zip(SET_BY_MAYANK, (c["work_order"], c["company_id"], c["company_name"], c["market"], c["arm"])):
            ws.cell(row=next_row, column=columns[name], value=value)
        if c["arm"] == ARM_AGENT:
            write_agent_row(ws, columns, next_row, agent_rows[c["company_id"]])
        next_row += 1
    wb.save(out)

    print(f"RELEASE {out.name}  built on {'the blank template' if base_is_template else base_path.name}\n")
    passed, dirty, kept_fonts = verify(out, base_path, base_is_template, expected, draw, agent_rows, restored, rules_rows)
    gate_ok = gate(out)
    if not (passed and gate_ok):
        failed = out.with_name(out.stem + ".FAILED.xlsx")
        out.rename(failed)
        raise SystemExit(f"release failed its checks; kept for inspection as {failed.name}, not to be sent")
    entry = {
        "file": out.name, "sha256": sha256(out), "built_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "to": "SDR", "built_on": "blank template v3" if base_is_template else base_path.name,
        "through_order": args.through_order, "rows": len(expected),
        "new_rows": sum(c["_new"] for c in expected),
        "sdr_alone_rows": sum(c["arm"] == ARM_SDR for c in expected),
        "with_agent_rows": sum(c["arm"] == ARM_AGENT for c in expected),
        "agent_rows_on_sdr_alone_companies": dirty,
        "grey_cells_restored": len(restored),
        "sdr_cells_kept_in_another_font": kept_fonts,
        "rules_added_to_instructions": len(rules_rows) > 0,
        "agent_runs": [Path(folder).name for folder in args.agent_runs],
        "agent_rows": [agent_row_record(c, args.agent_runs) for c in expected if c["_new"] and c["arm"] == ARM_AGENT],
    }
    (data / "master-log").mkdir(parents=True, exist_ok=True)
    with open(data / "master-log" / "releases.jsonl", "a", encoding="utf-8") as handle:
        handle.write(json.dumps(entry) + "\n")
    print(f"\nLOGGED in master-log/releases.jsonl: {entry}")
    print(f'{"SEALED" if seal(out) else "NOT SEALED"}  {out.name} is read-only on disk')


if __name__ == "__main__":
    main()
