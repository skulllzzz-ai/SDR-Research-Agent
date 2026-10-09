"""Read, check and file the lead's scored return (the lead's scoring workbook, sent back after a lead release).

    python lead_return.py file --received "<name as received>" --filed lead-YYYY-MM-DD-in.xlsx --received-on YYYY-MM-DD \
        [--note "..."]
    python lead_return.py check <file.xlsx> --answers lead-YYYY-MM-DD-out.xlsx

The lead's file is kept in data-exports/lead/in/ under the name it came with, and a byte-identical copy goes in under
the contract name (cp -n) before `file` runs. `file` checks the two are identical, reads the copy and appends one
record to data-exports/master-log/lead-returns.jsonl: the release it answers (the oldest lead release in force that
has no return yet), its hash and what the read found. A filed return is never edited.

The read (build_master_log.py uses it too) checks that the lead's scores survived and that nothing else moved:
  - the scoring layout is intact, and every row of the release is there, with its code, in the same order;
  - the row code, company, market and every research cell equal the release as sent. The release is read from the
    newest copy of it whose hash still matches the release log: a re-issue keeps rows, codes, order and values, so
    the copy it re-issued can stand in when the re-issued file itself was later saved over;
  - Acceptable, Correctness and Completeness hold a value from their list, the two stamps are times, the end not
    before the start.
Lead scoring minutes are whole minutes between the two stamps (Ctrl+Shift+; stamps minutes only, so 0 means scored
within the minute). Aggregates only on screen: the comments go to the master log and are never printed.
"""

import argparse
import hashlib
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from build_lead_view import (
    HEADERS, LISTS, RECHECK_HEADERS, RECHECK_QUESTION, RECHECK_SHEET, SCORING, SHEET, SHORT_ROW, clean, exchange_files, header_row,
    lead_releases, logged_releases,
)
from contract import RESEARCH_FIELDS
from sdr_return import as_time

HERE = Path(__file__).resolve().parent
VIEW = ["Row code", "Company name", "Market"] + RESEARCH_FIELDS + SHORT_ROW
RECHECK_VIEW = ["Row code", "Company name", "Market"] + SHORT_ROW
SCORES = list(LISTS)


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def intact_copy(data, name):
    """(path, note) of the newest copy of release `name` whose hash matches the release log, or (None, note)."""
    entries = {entry["file"]: entry for entry in logged_releases(data) if entry.get("to") == "Lead"}
    notes = []
    while name:
        entry, path = entries.get(name), Path(data) / "lead" / "out" / name
        if entry is None or not path.exists():
            return None, "; ".join(notes + [f"{name} is not in the release log or not on disk"])
        if sha256(path) == entry["sha256"]:
            return path, "; ".join(notes) or None
        notes.append(f"{name} no longer matches its logged hash")
        name = entry.get("reissue_of")
        if name:
            notes.append(f"checked against {name}, which it re-issued with the same rows, codes, order and values")
    return None, "; ".join(notes)


def table(ws):
    """(header row, column of each heading, [(sheet row, row code)])."""
    h = header_row(ws)
    column = {ws.cell(row=h, column=c).value: c for c in range(1, ws.max_column + 1) if ws.cell(row=h, column=c).value}
    rows = [(r, clean(ws.cell(row=r, column=1).value)) for r in range(h + 1, ws.max_row + 1) if clean(ws.cell(row=r, column=1).value)]
    return h, column, rows


def read_recheck(wb, out, key, data, answers):
    """A returned dead-end re-check (DECISIONS.md 18): one Yes or No per dead-end row and a comment, checked like a scoring
    return: the layout, every row with its code in order, the shown cells unchanged against the release, listed answers."""
    checks = out["checks"]
    out["kind"] = "dead-end re-check"
    ws = wb[RECHECK_SHEET]
    found = [r for r in range(1, 61) if [ws.cell(row=r, column=c).value for c in range(1, len(RECHECK_HEADERS) + 1)] == RECHECK_HEADERS]
    checks["re-check layout intact (headers)"] = bool(found)
    if not found:
        return out
    h = found[0]
    rows = [(r, clean(ws.cell(row=r, column=1).value)) for r in range(h + 1, ws.max_row + 1) if clean(ws.cell(row=r, column=1).value)]
    reference, out["reference_note"] = intact_copy(data, answers)
    out["checked_against"] = reference.name if reference else None
    checks["the release is on disk as logged, or a copy that stands in for it"] = reference is not None
    if reference:
        ref = load_workbook(reference)[RECHECK_SHEET]
        ref_rows = [(r, clean(ref.cell(row=r, column=1).value)) for r in range(h + 1, ref.max_row + 1) if clean(ref.cell(row=r, column=1).value)]
        checks["every row of the release is there, with its code, in the same order"] = [c for _, c in rows] == [c for _, c in ref_rows]
        changed = Counter(name for (r, _), (rr, _) in zip(rows, ref_rows) for c, name in enumerate(RECHECK_VIEW, start=1)
                          if clean(ws.cell(row=r, column=c).value) != clean(ref.cell(row=rr, column=c).value))
        out["cells_changed_by_column"] = dict(changed)
        checks["row code, company, market and the dead-end cells equal the release"] = not changed
    checks["every row code is in the lead key"] = all(code in key for _, code in rows)
    last = rows[-1][0] if rows else h
    outside = sum(1 for row in ws.iter_rows() for c in row if c.value not in (None, "") and (c.column > len(RECHECK_HEADERS) or c.row > last))
    checks["nothing typed outside the table"] = outside == 0
    q, comment_col = RECHECK_HEADERS.index(RECHECK_QUESTION) + 1, RECHECK_HEADERS.index("Comment") + 1
    for r, code in rows:
        answer, comment = clean(ws.cell(row=r, column=q).value), clean(ws.cell(row=r, column=comment_col).value)
        if answer and answer not in ("Yes", "No"):
            out["issues"]["answer: not Yes or No"] += 1
        entry = key.get(code, {})
        out["rows"].append({"row_code": code, "company_id": entry.get("company_id"), "row_type": entry.get("row_type"),
                            "source": entry.get("source"), "scored": False, "recheck": answer or None, "comment": comment or None})
    checks["answers are Yes or No"] = not out["issues"]
    answered = [r for r in out["rows"] if r["recheck"]]
    out["summary"] = {"rows": len(out["rows"]), "scored_rows": 0, "answered_rows": len(answered),
                      "answered_rows_by_type": dict(Counter(r["row_type"] for r in answered)),
                      "scored_rows_by_type": {}, "answers": dict(Counter(r["recheck"] for r in answered)),
                      "comments": sum(r["comment"] is not None for r in out["rows"])}
    return out


def read_return(path, data, answers):
    """Every row of a lead return with its scores and scoring minutes, and the checks that the scores survived."""
    data = Path(data)
    key = {e["row_code"]: e for e in json.loads((data / "master-log" / "lead-key.json").read_text(encoding="utf-8"))["rows"]}
    out = {"file": Path(path).name, "sha256": sha256(path), "answers": answers, "checks": {}, "rows": [], "issues": Counter()}
    checks = out["checks"]
    wb = load_workbook(path)
    if RECHECK_SHEET in wb.sheetnames:
        return read_recheck(wb, out, key, data, answers)
    out["kind"] = "scoring"
    checks["scoring sheet present"] = SHEET in wb.sheetnames
    if not checks["scoring sheet present"]:
        return out
    ws = wb[SHEET]
    h, column, rows = table(ws)
    checks["scoring layout intact (headers)"] = [ws.cell(row=h, column=c).value for c in range(1, len(HEADERS) + 1)] == HEADERS
    if not checks["scoring layout intact (headers)"]:
        return out
    reference, out["reference_note"] = intact_copy(data, answers)
    out["checked_against"] = reference.name if reference else None
    checks["the release is on disk as logged, or a copy that stands in for it"] = reference is not None
    if reference:
        ref_ws = load_workbook(reference)[SHEET]
        _, ref_column, ref_rows = table(ref_ws)
        checks["every row of the release is there, with its code, in the same order"] = [c for _, c in rows] == [c for _, c in ref_rows]
        changed = Counter()
        for (r, _), (rr, _) in zip(rows, ref_rows):
            for name in VIEW:
                if clean(ws.cell(row=r, column=column[name]).value) != clean(ref_ws.cell(row=rr, column=ref_column[name]).value):
                    changed[name] += 1
        out["cells_changed_by_column"] = dict(changed)
        checks["row code, company, market and research cells equal the release"] = not changed
    checks["every row code is in the lead key"] = all(code in key for _, code in rows)
    last = rows[-1][0] if rows else h
    outside = sum(1 for row in ws.iter_rows() for c in row if c.value not in (None, "") and (c.column > len(HEADERS) or c.row > last))
    checks["nothing typed outside the table"] = outside == 0

    for r, code in rows:
        values = {name: clean(ws.cell(row=r, column=column[name]).value) for name in SCORES}
        raw = {name: ws.cell(row=r, column=column[name]).value for name in ("Scoring start", "Scoring end")}
        start, end = as_time(raw["Scoring start"]), as_time(raw["Scoring end"])
        comment = clean(ws.cell(row=r, column=column["Comment"]).value)
        scored = all(values.values())
        if any(values.values()) and not scored:
            out["issues"]["scores: some but not all three filled"] += 1
        for name in SCORES:
            if values[name] and values[name] not in LISTS[name]:
                out["issues"][f"{name}: not a listed value"] += 1
        for name, value in raw.items():
            if value not in (None, "") and as_time(value) is None:
                out["issues"][f"{name}: not a time"] += 1
        minutes = None
        if start and end:
            delta = (datetime.combine(date.min, end) - datetime.combine(date.min, start)).total_seconds() / 60
            if delta < 0:
                out["issues"]["Scoring end before Scoring start"] += 1
            else:
                minutes = int(delta)
        elif scored:
            out["issues"]["scored without both stamps"] += 1
        entry = key.get(code, {})
        out["rows"].append({"row_code": code, "company_id": entry.get("company_id"), "row_type": entry.get("row_type"),
                            "source": entry.get("source"), "scored": scored, **{n.lower(): values[n] or None for n in SCORES},
                            "start": start, "end": end, "minutes": minutes, "comment": comment or None})
    checks["scores hold listed values; stamps are times, end not before start"] = not out["issues"]
    scored = [r for r in out["rows"] if r["scored"]]
    out["summary"] = {
        "rows": len(out["rows"]), "scored_rows": len(scored),
        "scored_rows_by_type": dict(Counter(r["row_type"] for r in scored)),
        "unscored_rows": len(out["rows"]) - len(scored),
        **{n.lower(): dict(Counter(r[n.lower()] for r in scored)) for n in SCORES},
        "rows_with_both_stamps": sum(r["minutes"] is not None for r in out["rows"]),
        "lead_minutes_total": sum(r["minutes"] for r in out["rows"] if r["minutes"] is not None),
        "comments": sum(r["comment"] is not None for r in out["rows"]),
    }
    return out


def report(result):
    """Aggregates only: counts and check names, never a company name or a comment."""
    print(f'FILE  {result["file"]}  sha256 {result["sha256"][:12]}…  answers {result["answers"]}')
    if result.get("checked_against"):
        print(f'      checked against {result["checked_against"]}' + (f' ({result["reference_note"]})' if result.get("reference_note") else ""))
    for name, passed in result["checks"].items():
        print(f'{"PASS" if passed else "FAIL"}  {name}')
    if result.get("cells_changed_by_column"):
        print(f'      cells changed, by column: {result["cells_changed_by_column"]}')
    for label, n in sorted(result["issues"].items(), key=lambda kv: (-kv[1], kv[0])):
        print(f"      {n} row(s): {label}")
    s = result.get("summary")
    if s and result.get("kind") == "dead-end re-check":
        print(f'ROWS  {s["rows"]} dead-end rows, {s["answered_rows"]} answered ({s["answered_rows_by_type"]}); answers {s["answers"]}; '
              f'{s["comments"]} comments')
    elif s:
        print(f'ROWS  {s["rows"]} rows, {s["scored_rows"]} scored ({s["scored_rows_by_type"]}), {s["unscored_rows"]} unscored; '
              f'stamps on {s["rows_with_both_stamps"]}, {s["lead_minutes_total"]} lead minutes; {s["comments"]} comments')
        for name in SCORES:
            print(f'      {name}: {s[name.lower()]}')


def file_return(data, received, filed, received_on, note=None):
    """Checks the two copies are identical, reads the filed copy against the release it answers, records the filing."""
    folder = data / "lead" / "in"
    filed_path, received_path = folder / filed, folder / received
    assert filed_path.exists(), f"{filed} must already be in {folder} (a byte-identical copy of the file as received)"
    assert received_path.exists(), f"{received} must be in {folder}, kept as received"
    assert sha256(received_path) == sha256(filed_path), "the filed copy is not byte-identical to the received file"
    assert filed_path in exchange_files(folder, "in"), f"{filed} is not a contract name (lead-YYYY-MM-DD-in.xlsx)"
    log = data / "master-log" / "lead-returns.jsonl"
    before = [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line.strip()] if log.exists() else []
    if any(entry["sha256"] == sha256(filed_path) and entry["filed_as"] == filed for entry in before):
        raise SystemExit(f"{filed} is already filed (same name and hash). Nothing written.")
    releases = [p.name for p in lead_releases(data)]
    answers = releases[len(before)] if len(before) < len(releases) else None
    assert answers, "every lead release in force already has a return: nothing for this file to answer"
    result = read_return(filed_path, data, answers)
    report(result)
    record = {"kind": "Lead return", "filed_as": filed, "received_as": received, "sha256": result["sha256"], "received_on": received_on,
              "filed_at": datetime.now().astimezone().isoformat(timespec="seconds"), "answers": answers, "from": "Lead",
              "return_of": result.get("kind"), "scored_rows": result.get("summary", {}).get("scored_rows", 0),
              "answered_rows": result.get("summary", {}).get("answered_rows"), "checked_against": result.get("checked_against"),
              "reference_note": result.get("reference_note"), "checks": result["checks"], "issue_counts": dict(result["issues"]),
              "summary": result.get("summary"), "note": note}
    with open(log, "a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    print(f"FILED {filed} (received as {received!r} on {received_on}), answers {answers}: recorded in {log.name}")
    if not all(result["checks"].values()):
        raise SystemExit("filed as received, but a check failed: the scores are not ingested until it is understood")
    return record


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("mode", choices=["file", "check"])
    parser.add_argument("path", nargs="?", help="check: the workbook to read")
    parser.add_argument("--answers", help="check: the lead release the workbook answers")
    parser.add_argument("--received", help="file: the name the file came with, as kept in lead/in/")
    parser.add_argument("--filed", help="file: the contract name, lead-YYYY-MM-DD-in.xlsx, already copied into lead/in/")
    parser.add_argument("--received-on", help="file: the date it was received, YYYY-MM-DD")
    parser.add_argument("--note", help="file: anything the record should carry (who, when, what was odd)")
    parser.add_argument("--data-root", default=str(HERE / "data-exports"))
    args = parser.parse_args()
    data = Path(args.data_root)
    if args.mode == "check":
        assert args.path and args.answers, "check needs the workbook and --answers"
        result = read_return(args.path, data, args.answers)
        report(result)
        sys.exit(0 if all(result["checks"].values()) else 1)
    assert args.received and args.filed and args.received_on, "file needs --received, --filed and --received-on"
    file_return(data, args.received, args.filed, args.received_on, note=args.note)


if __name__ == "__main__":
    main()
