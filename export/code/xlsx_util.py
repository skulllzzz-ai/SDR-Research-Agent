"""Shared workbook helpers: house styles, never-overwrite file names, and the xlsx gate.

The gate is the workspace's spreadsheet gate (shared/deliverable-gates.md): every XML part must
load in a real XML parser, sheet dimensions must tie to the cells, and the file must open in Excel
through COM without repair (excel_open_check.ps1). Run it on any workbook:

    python xlsx_util.py <file.xlsx> [more.xlsx ...]
"""

import os
import re
import stat
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries

HERE = Path(__file__).resolve().parent

FONT = "Arial"
BASE = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=13, bold=True)
HEADER_FILL = PatternFill("solid", start_color="D9E1F2")
WARN_FILL = PatternFill("solid", start_color="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
# Agent-row colours, the same three the template's Lists sheet shows the SDR.
SURE_FILL = PatternFill("solid", start_color="C6EFCE")
GUESSED_FILL = PatternFill("solid", start_color="FFEB9C")
NOT_FOUND_FILL = PatternFill("solid", start_color="FFC7CE")


def seal(path):
    """Marks a sent file read-only on disk once it is built and logged (Mayank, 9 Oct: the lead's scores had been saved
    over a sent copy). Excel then opens it read-only and asks for a new name when someone saves. True when sealed."""
    os.chmod(path, stat.S_IREAD)
    return not os.access(path, os.W_OK)


def next_version_path(folder, stem, ext=".xlsx"):
    """folder/stem.ext, or stem-2.ext, stem-3.ext ... so no version is ever overwritten."""
    folder = Path(folder)
    candidate = folder / f"{stem}{ext}"
    n = 1
    while candidate.exists():
        n += 1
        candidate = folder / f"{stem}-{n}{ext}"
    return candidate


def write_table(ws, headers, rows, widths=None, start_row=1):
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=name)
        cell.font, cell.fill, cell.border, cell.alignment = BOLD, HEADER_FILL, BOX, WRAP_TOP
        if widths:
            ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    for r, row in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font, cell.alignment = BASE, WRAP_TOP
    return start_row + len(rows)


def xml_gate(path):
    """Parse every part with a real XML parser and tie each sheet's dimension to its cells."""
    problems, parts, sheets = [], 0, {}
    with zipfile.ZipFile(path) as package:
        bad = package.testzip()
        if bad:
            problems.append(f"zip entry fails its checksum: {bad}")
        for name in package.namelist():
            if "\\" in name:
                problems.append(f"backslash in entry name: {name}")
            if not name.endswith((".xml", ".rels", ".vml")):
                continue
            parts += 1
            data = package.read(name)
            if not data.lstrip().startswith(b"<"):
                problems.append(f"{name} does not start with markup")
            try:
                root = ElementTree.fromstring(data)
            except ElementTree.ParseError as error:
                problems.append(f"{name} does not parse: {error}")
                continue
            if name.startswith("xl/worksheets/sheet"):
                ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                dimension = root.find("m:dimension", ns)
                refs = [c.get("r") for c in root.iterfind("m:sheetData/m:row/m:c", ns)]
                max_row = max((int(re.sub(r"[A-Z]+", "", r)) for r in refs), default=1)
                declared = dimension.get("ref") if dimension is not None else None
                sheets[name] = {"dimension": declared, "cells": len(refs), "last_row": max_row}
                if declared:
                    _, _, _, declared_last_row = range_boundaries(declared.upper())
                    if declared_last_row != max_row:
                        problems.append(f"{name}: dimension {declared} but cells reach row {max_row}")
    return {"parts": parts, "sheets": sheets, "problems": problems}


def excel_open(path):
    """Open the file in Excel through COM. Excel does not attempt recovery from the object model,
    so a file that needs repair fails here. Returns (opened, message)."""
    script = HERE / "excel_open_check.ps1"
    done = subprocess.run(
        ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), "-Path", str(Path(path).resolve())],
        capture_output=True, text=True, timeout=180,
    )
    message = (done.stdout or done.stderr).strip()
    return done.returncode == 0 and message.startswith("OPENED"), message


def gate(path):
    result = xml_gate(path)
    opened, message = excel_open(path)
    ok = not result["problems"] and opened
    print(f'{"GATE PASS" if ok else "GATE FAIL"}  {Path(path).name}')
    print(f'  XML parts parsed: {result["parts"]}, problems: {result["problems"] or "none"}')
    for name, info in result["sheets"].items():
        print(f'  {name}: dimension {info["dimension"]}, {info["cells"]} cells, last row {info["last_row"]}')
    print(f"  Excel: {message}")
    return ok


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(0 if all([gate(p) for p in sys.argv[1:]]) else 1)
