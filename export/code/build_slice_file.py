"""Build the slice file: one With-agent row, for one company outside the 30, for the SDR to check
as a rehearsal (spec exit condition 8). Same blank template, same instructions and rules, same
marker colours, source comments and flag band as a real release, so the rehearsal is the real thing.

    python build_slice_file.py --run data-exports/smoke-slice/<slice folder>/<company id> [--date YYYY-MM-DD]

Writes data-exports/smoke-slice/slice-YYYY-MM-DD-out.xlsx and never overwrites one. The SDR's return
is filed beside it as slice-YYYY-MM-DD-in.xlsx (kept as received, like every return) and read with
    python sdr_return.py check data-exports/smoke-slice/slice-YYYY-MM-DD-in.xlsx --draw none
whose Step 1 minutes are the SDR's checking time the exit condition asks for.
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_sdr_release as release
from contract import (
    AGENT_FLAG, ARM_AGENT, FIELD_KEYS, FIRST_DATA_ROW, HEADER_ROW, RESEARCH_COLUMNS, SET_BY_MAYANK,
    SHEET_INSTRUCTIONS, SHEET_LISTS, SHEET_RESEARCH, SHEET_SNAPSHOT, TIME_STAMPS, VERDICT_COLUMNS,
)
from validate_row import validate
from xlsx_util import gate, seal

HERE = Path(__file__).resolve().parent


def build(run_dir, out):
    run = json.loads((run_dir / "run.json").read_text(encoding="utf-8"))
    row = json.loads((run_dir / "row.json").read_text(encoding="utf-8"))
    company = {k: run[k] for k in ("company_id", "company_name", "market")}
    problems = validate(row, company)
    assert not problems, f"the slice row does not pass the validator: {problems}"
    assert not run.get("probe_not_an_agent_row"), "the slice row must come from the pinned agent, not a probe"
    wb = load_workbook(release.TEMPLATE)
    ws = wb[SHEET_RESEARCH]
    columns = release.header_columns(ws)
    release.clear_example_row(ws, columns)
    release.fix_instructions(wb[SHEET_INSTRUCTIONS])
    release.add_rules(wb[SHEET_INSTRUCTIONS])
    for name, value in zip(SET_BY_MAYANK, (1, company["company_id"], company["company_name"], company["market"], ARM_AGENT)):
        ws.cell(row=FIRST_DATA_ROW, column=columns[name], value=value)
    release.write_agent_row(ws, columns, FIRST_DATA_ROW, row)
    wb.save(out)
    return company, row, run


def verify(out, company, row):
    wb, blank = load_workbook(out), load_workbook(release.TEMPLATE)
    ws = wb[SHEET_RESEARCH]
    columns = release.header_columns(ws)
    results = []

    def check(name, passed, detail=""):
        results.append(bool(passed))
        print(f'{"PASS" if passed else "FAIL"}  {name}{"  [" + str(detail) + "]" if detail != "" else ""}')

    check("headers are the contract's 46 columns",
          [ws.cell(row=HEADER_ROW, column=c).value for c in range(1, len(RESEARCH_COLUMNS) + 1)] == RESEARCH_COLUMNS)
    rows = release.data_rows(ws, columns)
    check("exactly one company row, on row 3", rows == [(FIRST_DATA_ROW, company["company_id"])], rows)
    r = FIRST_DATA_ROW
    grey = tuple(ws.cell(row=r, column=columns[name]).value for name in SET_BY_MAYANK)
    check("grey cells: order 1, the slice company, its market, With agent",
          grey == (1, company["company_id"], company["company_name"], company["market"], ARM_AGENT))
    faithful = coloured = commented = 0
    for name, key in FIELD_KEYS.items():
        entry, cell = row["fields"][key], ws.cell(row=r, column=columns[name])
        faithful += cell.value == (entry["value"].strip() or None)
        coloured += release.marker_colour(cell) == release.MARKER_COLOURS[entry["marker"]]
        commented += (cell.comment is not None) == bool(entry["source"].strip()) and \
            (cell.comment is None or entry["source"].strip() in cell.comment.text)
    check("25 research cells equal the validated agent row", faithful == len(FIELD_KEYS), faithful)
    check("25 research cells carry their marker colour", coloured == len(FIELD_KEYS), coloured)
    check("a source comment exactly where the agent gave a source", commented == len(FIELD_KEYS), commented)
    flag = row["flag"]
    band = tuple(ws.cell(row=r, column=columns[name]).value for name in AGENT_FLAG)
    check("the flag band holds the agent's flag, reason and source",
          band == ("Yes" if flag["likely_dead_end"] == "yes" else "No", flag["reason"].strip() or None, flag["source"].strip() or None))
    left = [name for name in ["Status", "Date started"] + TIME_STAMPS + VERDICT_COLUMNS + ["Notes"]
            if ws.cell(row=r, column=columns[name]).value not in (None, "")]
    check("status, stamps, verdict, CRM and notes left empty for the SDR", not left, left)
    lines = [c.value for c in wb[SHEET_INSTRUCTIONS]["A"]]
    check("the rules added on 5 Oct are printed once", lines.count(release.RULES_ADDED[1]) == 1
          and all(line in lines for line in release.RULES_ADDED[1:]))
    check("the template's two stale instruction lines are replaced",
          all(old not in lines for old in release.INSTRUCTION_FIXES) and all(new in lines for new in release.INSTRUCTION_FIXES.values()))
    text = " ".join(str(c.value) for row_ in ws.iter_rows() for c in row_ if c.value is not None)
    check("the template's example row is gone", "EXAMPLE ROW" not in text and "EX00" not in text)
    for title in (SHEET_SNAPSHOT, SHEET_LISTS):
        check(f"{title}: as in the blank template", release.sheet_values(wb[title]) == release.sheet_values(blank[title]))
    print(f"\n{sum(results)} of {len(results)} checks passed")
    return all(results)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--run", required=True, help="the slice company's run folder, holding run.json and row.json")
    parser.add_argument("--date", default=datetime.now().strftime("%Y-%m-%d"))
    parser.add_argument("--out", help="for tests only: another file to write")
    args = parser.parse_args()
    out = Path(args.out) if args.out else HERE / "data-exports" / "smoke-slice" / f"slice-{args.date}-out.xlsx"
    if out.exists():
        raise SystemExit(f"{out.name} exists. A file sent out is never overwritten.")
    company, row, run = build(Path(args.run), out)
    print(f"SLICE FILE {out.name}: {company['company_id']} ({company['market']}), agent row of {run['agent_end']}, "
          f"{run['agent_minutes']} min, model {run['model_asked']} at {run['effort']}\n")
    ok = verify(out, company, row) and gate(out)
    if not ok:
        failed = out.with_name(out.stem + ".FAILED.xlsx")
        out.rename(failed)
        raise SystemExit(f"the slice file failed its checks; kept as {failed.name}, not to be sent")
    print(f'{"SEALED" if seal(out) else "NOT SEALED"}  {out.name} is read-only on disk')


if __name__ == "__main__":
    main()
