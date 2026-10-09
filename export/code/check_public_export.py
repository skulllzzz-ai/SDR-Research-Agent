"""The contact-and-SDR scrub gate for the public export (DECISIONS.md 18). Exit 0 only when every check passes.

    python check_public_export.py data-exports/public-export --draw data-exports/master-log/draw-2026-10-01.json
    python check_public_export.py --self-test --draw data-exports/master-log/draw-2026-10-01.json \
        --agent-runs data-exports/agent-runs/batch-v1 data-exports/agent-runs/batch-v2 \
        --smoke data-exports/smoke-slice/smoke-v1web-2026-10-06 data-exports/smoke-slice/smoke-v2-2026-10-07

It replaces the name scrub of the commit gate for this folder: the companies are public here, so what must not leak is a
contact detail or anything of the SDR's own. Over every file of the export:
  1. no email address (the documentation domains example.com, .org, .net of the made-up test values in the code excepted);
  2. no phone number (the two all-zero made-up numbers in validate_row.py's examples excepted);
  3. the CRM status is masked in both tables, and no CRM value the agent or the SDR wrote stands anywhere;
  4. no SDR stamp: the tables carry only their allowed columns and fields, none of them a time;
  5. no SDR research cell: none of the SDR's own texts (a cell of 12 characters or more that the agent did not write for
     that company) appears;
  6. no lead comment: no six-word run of any of the lead's comments appears;
  7. no reserve company, no line of the competitor block; competitor names only where the agent itself wrote them;
  8. only the allowed files: no workbook, no stream, no assembled prompt, no key or log of the exchange.
--self-test is the negative control: it plants an email, a phone, a CRM value, an SDR stamp, an SDR research cell and a
lead comment into scratch copies of the inputs, builds the export from them, and proves each appears 0 times in it, after
a positive control proves the same scan finds each in the inputs and in a deliberately leaky copy of the export.
"""

import argparse
import csv
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import time
from pathlib import Path

from openpyxl import load_workbook

from build_public_export import COMPANY_COLUMNS, EMAIL, MASKED, NOT_A_PHONE, PHONE, ROW_COLUMNS
from contract import FIELD_KEYS, RESEARCH_FIELDS, ROW_AGENT_ALONE
from results import newest_master_log, read_rows

HERE = Path(__file__).resolve().parent
ALLOWED_EMAIL_DOMAINS = ("example.com", "example.org", "example.net")
ALLOWED_PHONE_DIGITS = {"10000000000", "10000000001"}
TOP_FILES = {"README.md", "RESULTS.md", "DECISIONS.md", "BUILD-LOG.md", "companies.csv", "agent-rows.csv", "EXPORT-MANIFEST.json"}
CODE_SUFFIXES = {".py", ".json", ".ps1", ".md"}
ROW_FIELDS = RESEARCH_FIELDS + ["Flag: likely dead end", "Flag reason", "Verdict", "Dead-end kind", "Verdict reason", "CRM status"]
CRM_WORDS = ["Not attempted", "out of declared scope"]
# Built at run time, so this file, which the export carries in code/, holds none of them as written: the email and phone
# are split, the tokens joined, the comment's words kept in reverse order. The stamp is picked once the baseline is read.
PLANT = {"email": "q9.planted" + "@" + "scrubcheck-mail.com", "phone": " ".join(["+971", "4", "386", "2917"]),
         "crm": "-".join(["PLANTED", "CRM", "STATUS", "Q9"]),
         "sdr_cell": "-".join(["PLANTED", "SDR", "RESEARCH", "CELL", "Q9"]) + " " + " ".join(["sells", "only", "through", "its", "own", "boutiques"]),
         "comment": " ".join(reversed(["look", "second", "a", "needs", "row", "this", "says", "nine", "q", "comment", "lead", "planted"])),
         "stamp": None}
STAMP_CANDIDATES = [time(5, 43), time(4, 29), time(6, 17), time(3, 51), time(2, 37)]


def export_texts(folder):
    out = {}
    for p in sorted(Path(folder).rglob("*")):
        if p.is_file():
            out[str(p.relative_to(folder)).replace("\\", "/")] = p.read_bytes().decode("utf-8-sig", errors="replace")
    return out


def phones_in(text):
    found = []
    for m in PHONE.finditer(text):
        raw = m.group(0)
        digits = re.sub(r"\D", "", raw)
        if 7 <= len(digits) <= 15 and not NOT_A_PHONE.fullmatch(raw.replace("tel:", "").strip()) and digits not in ALLOWED_PHONE_DIGITS:
            found.append(raw)
    return found


def emails_in(text):
    return [m for m in EMAIL.findall(text) if not m.lower().rstrip(".").endswith(ALLOWED_EMAIL_DOMAINS)]


def six_word_runs(text):
    words = re.findall(r"[a-z']+", (text or "").lower())
    return {" ".join(words[i:i + 6]) for i in range(len(words) - 5)}


def stamp_forms(t):
    """How a planted stamp could show in a text, as patterns: 05:43 or 5:43 standing alone, and the Excel day fraction."""
    fraction = (t.hour * 60 + t.minute) / 1440
    return [rf"(?<![\d:]){t.hour:02d}:{t.minute:02d}(?![\d:])", rf"(?<![\d:]){t.hour}:{t.minute:02d}(?![\d:])",
            re.escape(f"{fraction:.6f}"[1:])]


def stamp_count(text, t):
    return sum(len(re.findall(form, text)) for form in stamp_forms(t))


def sdr_texts(master_rows, agent_text, public_names=()):
    """The SDR's own research texts: a cell of 12 characters or more on an SDR-alone or pair row that the agent never
    wrote (pair rows start from the agent's row, so a value the SDR kept is the agent's, and stays public). A company's
    own name, or a part of it, which the export shows anyway, is not the SDR's text."""
    names = {n.strip().lower() for n in public_names}
    out = set()
    for r in master_rows:
        if r["Row type"] == ROW_AGENT_ALONE:
            continue
        for name in RESEARCH_FIELDS + ["Dead-end reason with source", "Notes"]:
            value = r.get(name)
            if (isinstance(value, str) and len(value.strip()) >= 12 and value.strip() not in agent_text
                    and not any(value.strip().lower() in n for n in names)):
                out.add(value.strip())
    return out


def competitor_names(block):
    """The names in the competitor block: the pipe-separated list, and the group named before each colon."""
    names = set()
    for line in block.splitlines():
        if "|" in line:
            names.update(re.sub(r"\(.*?\)", "", part).strip() for part in line.split("|"))
        elif re.match(r"^[^:]{3,40}:\s", line) and not line.rstrip().endswith(":"):
            names.add(line.split(":")[0].strip())
    return sorted(n for n in names if len(n) > 3)


def check(folder, draw, master_rows, competitor_block):
    results = []

    def ok(name, passed, detail=""):
        results.append(bool(passed))
        print(f'{"PASS" if passed else "FAIL"}  {name}{"  [" + str(detail) + "]" if detail != "" else ""}')

    folder = Path(folder)
    texts = export_texts(folder)
    everything = "\n".join(texts.values())
    allowed = [f for f in texts if f in TOP_FILES or (f.startswith("code/") and Path(f).suffix in CODE_SUFFIXES)]
    ok("only the allowed files: the tables, the reports, the logs, the code and its prompts", len(allowed) == len(texts),
       sorted(set(texts) - set(allowed))[:5] or f"{len(texts)} files")
    ok("no workbook, stream, assembled prompt, lead key or exchange log", not any(
        f.endswith((".xlsx", ".jsonl")) or "assembled-prompt" in f or "lead-key" in f or f.endswith(("releases.jsonl", "returns.jsonl"))
        for f in texts))
    emails = {f: emails_in(t) for f, t in texts.items()}
    ok("no email address anywhere (documentation domains excepted)", not any(emails.values()),
       {f: len(v) for f, v in emails.items() if v} or 0)
    phones = {f: phones_in(t) for f, t in texts.items()}
    ok("no phone number anywhere (the made-up all-zero examples excepted)", not any(phones.values()),
       {f: len(v) for f, v in phones.items() if v} or 0)

    with open(folder / "companies.csv", encoding="utf-8-sig", newline="") as handle:
        companies = list(csv.reader(handle))
    with open(folder / "agent-rows.csv", encoding="utf-8-sig", newline="") as handle:
        agent_rows = list(csv.reader(handle))
    ok("the tables carry only their allowed columns (no stamp, no SDR minutes, no CRM value column)",
       companies[0] == COMPANY_COLUMNS and agent_rows[0] == ROW_COLUMNS)
    ok("agent-rows.csv holds only the agent's fields", {r[4] for r in agent_rows[1:]} <= set(ROW_FIELDS),
       sorted({r[4] for r in agent_rows[1:]} - set(ROW_FIELDS))[:3] or "")
    crm_col = COMPANY_COLUMNS.index("CRM status")
    crm_cells = [r[crm_col] for r in companies[1:]] + [r[5] for r in agent_rows[1:] if r[4] == "CRM status"]
    sdr_crm = {str(r["CRM status"]).strip() for r in master_rows if r["Row type"] != ROW_AGENT_ALONE and r.get("CRM status")}
    # the code names the CRM values as constants; the tables and reports must not carry any
    reports = "\n".join(t for f, t in texts.items() if f in ("companies.csv", "agent-rows.csv", "README.md", "RESULTS.md"))
    words_found = [w for w in CRM_WORDS + sorted(v for v in sdr_crm if len(v) > 8) if w in reports]
    ok("the CRM status is masked on every row, and no CRM value stands in the tables or reports",
       crm_cells and all(c == MASKED for c in crm_cells) and not words_found, words_found or f"{len(crm_cells)} masked cells")

    agent_text = "\n".join(r[5] for r in agent_rows[1:]) + "\n" + "\n".join(
        str(r.get(n) or "") for r in master_rows if r["Row type"] == ROW_AGENT_ALONE for n in RESEARCH_FIELDS)
    own = sdr_texts(master_rows, agent_text, [r[0] for r in companies[1:]])
    leaked = [t for t in own if t in everything]
    ok("no SDR research cell: none of the SDR's own texts appears", not leaked, f"{len(own)} SDR texts checked, {len(leaked)} found")
    comments = [str(r.get(k)) for r in master_rows for k in ("Lead comment", "Lead re-check comment") if r.get(k)]
    runs = set().union(*(six_word_runs(c) for c in comments)) if comments else set()
    found = runs & set().union(*(six_word_runs(t) for t in texts.values()))
    ok("no lead comment: no six-word run of the lead's comments appears", not found, f"{len(runs)} runs checked, {len(found)} found")

    reserves = sorted({r["company_name"] for r in draw["not_drawn"] + draw["pasted_reserves"]})
    low = everything.lower()
    hits = [n for n in reserves if len(n) > 3 and re.search(r"(?<![a-z])" + re.escape(n.lower()) + r"(?![a-z])", low)]
    ok("no reserve or not-drawn company appears", not hits, len(hits))
    block_lines = [line.strip() for line in competitor_block.splitlines() if len(line.strip()) >= 20]
    whole = [line for line in block_lines if line in everything]
    outside = [n for n in competitor_names(competitor_block)
               if any(re.search(r"(?<![a-z])" + re.escape(n.lower()) + r"(?![a-z])", t.lower()) for f, t in texts.items()
                      if f not in ("agent-rows.csv", "companies.csv"))]
    ok("no line of the competitor block appears; competitor names only where the agent itself wrote them",
       not whole and not outside, f"{len(whole)} lines, {len(outside)} names outside the agent's rows")
    prompts = [t for f, t in texts.items() if f.startswith("code/prompt/research-agent")]
    ok("the prompts keep the competitor placeholder", prompts and all("{{COMPETITOR_BLOCK}}" in t for t in prompts), f"{len(prompts)} prompts")
    print(f"\n{sum(results)} of {len(results)} gate checks passed")
    return all(results)


def scan_counts(text):
    """How many times each planted value shows in a text."""
    return {"email": text.count(PLANT["email"]), "phone": text.count(PLANT["phone"]), "crm": text.count(PLANT["crm"]),
            "stamp": stamp_count(text, PLANT["stamp"]), "sdr_cell": text.count(PLANT["sdr_cell"]),
            "comment": text.count(PLANT["comment"])}


def self_test(args):
    """The negative control on scratch copies: plant, build, prove 0; with a positive control on the inputs and a leaky copy."""
    root = Path(tempfile.mkdtemp(prefix="public-export-control-"))
    data = root / "data"
    (data / "master-log").mkdir(parents=True)
    master = Path(args.master_log) if args.master_log else newest_master_log(HERE / "data-exports" / "master-log")
    shutil.copyfile(master, data / "master-log" / master.name)
    runs = []
    for folder in args.agent_runs + args.smoke:
        target = data / "runs" / Path(folder).name
        for p in Path(folder).glob("*/row.json"):
            (target / p.parent.name).mkdir(parents=True, exist_ok=True)
            for name in ("row.json", "run.json"):
                shutil.copyfile(p.parent / name, target / p.parent.name / name)
        runs.append(target)
    v1, v2, smoke = runs[0], runs[1], runs[2:]
    baseline_out = root / "baseline"
    common = [sys.executable, str(HERE / "build_public_export.py"), "--draw", args.draw, "--data-root", str(data), "--no-gate",
              "--master-log", str(data / "master-log" / master.name), "--agent-runs", str(v1), str(v2), "--smoke", *map(str, smoke)]
    done = subprocess.run(common + ["--out", str(baseline_out)], capture_output=True, text=True, encoding="utf-8")
    baseline = "\n".join(export_texts(baseline_out).values()) if done.returncode == 0 else ""
    results = []

    def ok(name, passed, detail=""):
        results.append(bool(passed))
        print(f'{"PASS" if passed else "FAIL"}  {name}{"  [" + str(detail) + "]" if detail != "" else ""}')

    ok("baseline export built from the scratch copies", done.returncode == 0, done.stdout.strip().splitlines()[-1][:100] if done.stdout else done.stderr[-200:])
    PLANT["stamp"] = next(t for t in STAMP_CANDIDATES if stamp_count(baseline, t) == 0)
    ok("baseline: none of the planted values is there before planting", not any(scan_counts(baseline).values()),
       {**scan_counts(baseline), "stamp used": PLANT["stamp"].strftime("%H:%M")})

    # plant into the agent's row: emails and phones fields, a research note, a source link, the CRM status
    first = sorted(v1.glob("C*/row.json"))[0]
    row = json.loads(first.read_text(encoding="utf-8"))
    row["fields"]["emails"]["value"] = PLANT["email"]
    row["fields"]["phones"]["value"] = PLANT["phone"]
    row["fields"]["detailed_research"]["value"] += f" Write to {PLANT['email']} or call {PLANT['phone']}."
    row["fields"]["parent_company"]["source"] = f"mailto:{PLANT['email']}"
    row["crm_status"] = PLANT["crm"]
    first.write_text(json.dumps(row, ensure_ascii=False), encoding="utf-8")
    # plant into the master log: the SDR's CRM value, an SDR stamp, an SDR research cell, a lead comment
    path = data / "master-log" / master.name
    wb = load_workbook(path)
    ws = wb["Rows"]
    col = {c.value: i + 1 for i, c in enumerate(ws[1])}
    sdr_row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=col["Row type"]).value == "SDR alone")
    pair_row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=col["Row type"]).value == "Pair")
    ws.cell(row=sdr_row, column=col["CRM status"]).value = PLANT["crm"]
    ws.cell(row=sdr_row, column=col["Pass-1 start"]).value = PLANT["stamp"]
    ws.cell(row=sdr_row, column=col["Detailed research"]).value = PLANT["sdr_cell"]
    ws.cell(row=pair_row, column=col["Parent company"]).value = PLANT["sdr_cell"]
    ws.cell(row=pair_row, column=col["Check start"]).value = PLANT["stamp"]
    ws.cell(row=sdr_row, column=col["Lead comment"]).value = PLANT["comment"]
    wb.save(path)

    inputs = first.read_text(encoding="utf-8") + "\n" + "\n".join(
        c.value.strftime("%H:%M") if isinstance(c.value, time) else str(c.value)
        for row_ in load_workbook(path)["Rows"].iter_rows() for c in row_ if c.value is not None)
    found_in = scan_counts(inputs)
    ok("positive control: every planted value is in the inputs the export reads, and the scan finds it", all(found_in.values()), found_in)
    planted_out = root / "planted"
    done = subprocess.run(common + ["--out", str(planted_out)], capture_output=True, text=True, encoding="utf-8")
    ok("the export builds from the planted inputs", done.returncode == 0, done.stdout.strip().splitlines()[-1][:100] if done.stdout else done.stderr[-200:])
    exported = "\n".join(export_texts(planted_out).values()) if done.returncode == 0 else ""
    counts = scan_counts(exported)
    for key, n in counts.items():
        ok(f"negative control: the planted {key.replace('_', ' ')} appears 0 times in the export", n == 0, n)
    leaky = planted_out / "README.md"
    leaky.write_text(leaky.read_text(encoding="utf-8") + "\n" + " ".join(
        [PLANT["email"], PLANT["phone"], PLANT["crm"], PLANT["stamp"].strftime("%H:%M"), PLANT["sdr_cell"], PLANT["comment"]]), encoding="utf-8")
    leak = scan_counts("\n".join(export_texts(planted_out).values()))
    ok("positive control: the same scan finds every planted value in a deliberately leaky copy", all(leak.values()), leak)
    draw = json.loads(Path(args.draw).read_text(encoding="utf-8"))
    master_rows = read_rows(path)
    block = (HERE / "data-exports" / "list" / "competitor-block-2026-10-01.md").read_text(encoding="utf-8")
    print("\nThe gate on the leaky copy (it must fail):")
    caught = not check(planted_out, draw, master_rows, block)
    ok("the gate fails the leaky copy", caught)
    shutil.rmtree(root, ignore_errors=True)
    print(f"\n{sum(results)} of {len(results)} control checks passed")
    return all(results)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("folder", nargs="?")
    parser.add_argument("--draw", required=True)
    parser.add_argument("--master-log")
    parser.add_argument("--data-root", default=str(HERE / "data-exports"))
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--agent-runs", nargs=2)
    parser.add_argument("--smoke", nargs="*", default=[])
    args = parser.parse_args()
    if args.self_test:
        assert args.agent_runs, "--self-test needs --agent-runs <v1 batch> <v2 batch>"
        sys.exit(0 if self_test(args) else 1)
    draw = json.loads(Path(args.draw).read_text(encoding="utf-8"))
    master = Path(args.master_log) if args.master_log else newest_master_log(Path(args.data_root) / "master-log")
    block = (Path(args.data_root) / "list" / "competitor-block-2026-10-01.md").read_text(encoding="utf-8")
    print(f"CONTACT-AND-SDR SCRUB GATE  {args.folder}")
    sys.exit(0 if check(args.folder, draw, read_rows(master), block) else 1)


if __name__ == "__main__":
    main()
