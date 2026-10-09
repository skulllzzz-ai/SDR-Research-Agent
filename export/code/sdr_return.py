"""Read, check and file an SDR return (the SDR's workbook, sent back after a release).

    python sdr_return.py file --received "<name as received>" --filed sdr-YYYY-MM-DD-in.xlsx --received-on YYYY-MM-DD
    python sdr_return.py check <file.xlsx>
    python sdr_return.py --self-test

A return is filed in data-exports/sdr/in/ twice: as received (never touched) and as a byte-identical
copy under the contract name sdr-YYYY-MM-DD-in.xlsx, which the exchange scripts read. `file` checks
that the two are identical, reads the copy, prints aggregate checks and appends one record to
data-exports/master-log/returns.jsonl (a hash is never filed twice). build_master_log.py reads the
filed returns and files the SDR's Step 1 and Step 2 stamps as pass 1 and pass 2 on SDR-alone rows and
as check and completing on Pair rows.

Nothing the SDR typed is changed or "fixed" here. Issues are counted and carried next to the row.
"""

import argparse
import hashlib
import json
import sys
import tempfile
from collections import Counter
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

from contract import (
    AGENT_FLAG, ARM_AGENT, ARM_SDR, DEAD_END_ACTIVE, AGENT_DEAD_END_TYPES, FIRST_DATA_ROW, HEADER_ROW,
    PRIORITIES, PRODUCT_CATEGORIES, RESEARCH_COLUMNS, RESEARCH_FIELDS, SDR_UNSURE_VALUES, SET_BY_MAYANK,
    SHEET_INSTRUCTIONS, SHEET_LISTS, SHEET_RESEARCH, SHEET_SNAPSHOT, TIME_STAMPS, VERDICT_DEAD_END, VERDICT_PROSPECT,
)

HERE = Path(__file__).resolve().parent
STATUSES = ["Not started", "Step 1 done", "Complete"]
CRM_VALUES = ["New", "Active", "Dormant", "Not attempted"]
DEAD_END_TYPES = AGENT_DEAD_END_TYPES + [DEAD_END_ACTIVE]
STEP1, STEP2 = ("Step 1 start", "Step 1 end"), ("Step 2 start", "Step 2 end")
MGMT = ("Management asked at", "Management answer added at")


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def as_time(value):
    """A stamp cell as a time, or None. The SDR stamps hh:mm with Ctrl+Shift+;."""
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, str) and value.strip():
        for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p"):
            try:
                return datetime.strptime(value.strip(), fmt).time()
            except ValueError:
                continue
    return None


def minutes(start, end):
    """Whole minutes from start to end on one day; None when either is missing or end is earlier."""
    if start is None or end is None:
        return None
    delta = datetime.combine(date.min, end) - datetime.combine(date.min, start)
    return None if delta < timedelta(0) else int(delta.total_seconds() // 60)


def header_columns(ws):
    return {ws.cell(row=HEADER_ROW, column=c).value: c for c in range(1, ws.max_column + 1)
            if ws.cell(row=HEADER_ROW, column=c).value}


def sheet_rows(ws, columns):
    """Row number and the 46 values by header, for every data row holding a company id."""
    found = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        company_id = ws.cell(row=r, column=columns["Company ID"]).value
        if company_id not in (None, ""):
            found.append((r, {name: ws.cell(row=r, column=columns[name]).value for name in RESEARCH_COLUMNS}))
    return found


def read_return(path, draw=None, release_path=None):
    """Every company row of the RESEARCH sheet with its stamps, minutes, values and issues, the
    Pass 1 snapshot per row, and the file-level checks. Issues are short labels, so they can be
    counted across rows without naming a company."""
    wb = load_workbook(path)
    out = {"file": Path(path).name, "sha256": sha256(path), "rows": [], "file_checks": {}, "issue_counts": Counter()}
    checks = out["file_checks"]
    checks["sheets_present"] = all(s in wb.sheetnames for s in (SHEET_INSTRUCTIONS, SHEET_RESEARCH, SHEET_SNAPSHOT, SHEET_LISTS))
    if not checks["sheets_present"]:
        checks["sheets"] = wb.sheetnames
        return out
    ws, snap = wb[SHEET_RESEARCH], wb[SHEET_SNAPSHOT]
    columns = header_columns(ws)
    checks["headers_match_contract"] = [ws.cell(row=HEADER_ROW, column=c).value for c in range(1, len(RESEARCH_COLUMNS) + 1)] == RESEARCH_COLUMNS
    checks["snapshot_headers_match"] = [snap.cell(row=HEADER_ROW, column=c).value for c in range(1, len(RESEARCH_COLUMNS) + 1)] == RESEARCH_COLUMNS
    if not checks["headers_match_contract"]:
        return out
    cells = [c for row in ws.iter_rows() for c in row]
    checks["formulas"] = sum(1 for c in cells if isinstance(c.value, str) and c.value.startswith("="))
    checks["comments"] = sum(1 for c in cells if c.comment)
    checks["data_validations"] = len(ws.data_validations.dataValidation)
    if release_path:
        rel = load_workbook(release_path)
        for name, key in ((SHEET_INSTRUCTIONS, "instructions_unchanged"), (SHEET_LISTS, "lists_unchanged")):
            checks[key] = [[c.value for c in r] for r in wb[name].iter_rows()] == [[c.value for c in r] for r in rel[name].iter_rows()]
    by_id = {c["company_id"]: c for c in draw["companies"]} if draw else {}
    snapshots = {}
    if checks["snapshot_headers_match"]:
        for r, values in sheet_rows(snap, header_columns(snap)):
            snapshots.setdefault(str(values["Company ID"]), values)
    checks["snapshot_rows"] = len(snapshots)

    rows = []
    for r, v in sheet_rows(ws, columns):
        cid = str(v["Company ID"])
        issues = []
        row = {"sheet_row": r, "company_id": cid, "order": v["Order"], "company_name": v["Company name"], "market": v["Market"],
               "arm": v["Arm"], "status": v["Status"], "date_started": v["Date started"],
               "stamps": {name: as_time(v[name]) for name in TIME_STAMPS},
               "fields": {name: v[name] for name in RESEARCH_FIELDS}, "verdict": v["Verdict"], "dead_end_type": v["Dead-end type"],
               "dead_end_reason": v["Dead-end reason with source"], "crm_status": v["CRM status"], "notes": v["Notes"],
               "agent_flag": {name: v[name] for name in AGENT_FLAG}}
        for name in TIME_STAMPS:
            if v[name] not in (None, "") and row["stamps"][name] is None:
                issues.append(f"{name}: unreadable stamp")
        # the grey columns are Mayank's: compared with the draw
        if cid in by_id:
            c = by_id[cid]
            for name, want in zip(SET_BY_MAYANK, (c["work_order"], cid, c["company_name"], c["market"], c["arm"])):
                if v[name] != want:
                    issues.append(f"{name}: grey cell changed" + (" (blank)" if v[name] in (None, "") else ""))
        elif draw:
            issues.append("company id not in the draw")
        if v["Status"] not in STATUSES:
            issues.append("Status: not a listed value")
        complete = v["Status"] == "Complete"
        s = row["stamps"]
        s1, s2 = (s[STEP1[0]], s[STEP1[1]]), (s[STEP2[0]], s[STEP2[1]])
        stopped_dead_end = v["Verdict"] == VERDICT_DEAD_END and s2 == (None, None)
        row["pass1_minutes"], row["pass2_minutes"] = minutes(*s1), minutes(*s2)
        if complete or v["Status"] == "Step 1 done":
            if None in s1:
                issues.append("Step 1: stamp missing")
            elif row["pass1_minutes"] is None:
                issues.append("Step 1: end before start")
        if complete and not stopped_dead_end:
            if None in s2:
                issues.append("Step 2: stamp missing")
            elif row["pass2_minutes"] is None:
                issues.append("Step 2: end before start")
            elif s2 == s1:
                issues.append("Step 2: stamps equal Step 1 (one block, two passes)")
                row["pass2_minutes"] = None
            elif s1[1] is not None and s2[0] < s1[1]:
                issues.append("Step 2: starts before Step 1 ends")
        if (s[MGMT[0]] is None) != (s[MGMT[1]] is None):
            issues.append("Management: one of the two stamps missing")
        # values
        if complete:
            if v["Verdict"] not in (VERDICT_PROSPECT, VERDICT_DEAD_END):
                issues.append("Verdict: missing or not a listed value")
            if v["CRM status"] not in CRM_VALUES:
                issues.append("CRM status: missing or not a listed value")
            if v["Verdict"] == VERDICT_DEAD_END:
                if v["Dead-end type"] not in DEAD_END_TYPES:
                    issues.append("Dead-end type: missing or not a listed value")
                if not v["Dead-end reason with source"]:
                    issues.append("Dead-end reason: missing")
            if v["Verdict"] == VERDICT_PROSPECT:
                if v["Dead-end type"]:
                    issues.append("Dead-end type on a Prospect row")
                if v["Priority"] not in PRIORITIES:
                    issues.append("Priority: not on the AAA to E scale")
                blanks = [name for name in RESEARCH_FIELDS if v[name] in (None, "")]
                issues += [f"{name}: blank on a Prospect row" for name in blanks]
            for name in PRODUCT_CATEGORIES:
                if v[name] in SDR_UNSURE_VALUES:
                    issues.append(f"{name}: unsure value, counted as not found (decision 2026-10-05)")
                elif v[name] not in ("Y", "N", None, ""):
                    issues.append(f"{name}: not Y or N")
            if v["Arm"] == ARM_SDR:
                snapshot = snapshots.get(cid)
                row["snapshot"] = snapshot is not None
                if snapshot is None:
                    issues.append("Pass 1 snapshot: missing")
                else:
                    row["fields_changed_in_pass_2"] = sum(1 for name in RESEARCH_FIELDS if snapshot[name] != v[name])
            if v["Arm"] == ARM_AGENT and any(v[name] not in (None, "") for name in AGENT_FLAG) is False:
                issues.append("Agent flag band empty on a With-agent row")
        if v["Arm"] == ARM_SDR and any(v[name] not in (None, "") for name in AGENT_FLAG):
            issues.append("Agent flag band filled on an SDR-alone row")
        row["issues"] = issues
        rows.append(row)
    # block stamps: the same Step 1 pair on several rows means the SDR stamped a block, not a company
    pairs = Counter((row["stamps"][STEP1[0]], row["stamps"][STEP1[1]]) for row in rows if None not in (row["stamps"][STEP1[0]], row["stamps"][STEP1[1]]))
    for row in rows:
        pair = (row["stamps"][STEP1[0]], row["stamps"][STEP1[1]])
        if None not in pair and pairs[pair] > 1:
            row["issues"].append(f"Step 1: stamps shared with {pairs[pair] - 1} other row(s)")
            row["pass1_minutes_shared"] = True
        row["sdr_minutes"], row["minutes_basis"] = sdr_minutes(row, pairs.get(pair, 1))
    for row in rows:
        out["issue_counts"].update(row["issues"])
    out["rows"] = rows
    with_minutes = [r["sdr_minutes"] for r in rows if r["sdr_minutes"] is not None]
    out["summary"] = {
        "rows": len(rows), "complete": sum(r["status"] == "Complete" for r in rows),
        "sdr_alone_rows": sum(r["arm"] == ARM_SDR for r in rows), "with_agent_rows": sum(r["arm"] == ARM_AGENT for r in rows),
        "prospects": sum(r["verdict"] == VERDICT_PROSPECT for r in rows), "dead_ends": sum(r["verdict"] == VERDICT_DEAD_END for r in rows),
        "rows_with_issues": sum(bool(r["issues"]) for r in rows), "snapshot_rows": len(snapshots),
        "pass1_minutes_readable": sum(r["pass1_minutes"] is not None and not r.get("pass1_minutes_shared") for r in rows),
        "pass2_minutes_readable": sum(r["pass2_minutes"] is not None for r in rows),
        "sdr_minutes_rows": len(with_minutes), "sdr_minutes_total": round(sum(with_minutes), 1),
    }
    return out


def sdr_minutes(row, shared_by):
    """The SDR's minutes on one row, both passes together, and how they were read. Decision 1 of
    2026-10-05 (DECISIONS.md): where one Step 1 span is stamped on several rows, each row gets the
    block average; where Step 2 equals Step 1, the one span covers both passes."""
    s = row["stamps"]
    s1, s2 = (s[STEP1[0]], s[STEP1[1]]), (s[STEP2[0]], s[STEP2[1]])
    span1, span2 = minutes(*s1), minutes(*s2)
    if span1 is None:
        return None, "unreadable: Step 1 missing or ends before it starts"
    if s2 == s1:
        total, basis = span1, "one span for both passes (Step 2 equals Step 1)"
    elif None in s2:
        total, basis = span1, "Step 1 only (Step 2 not stamped)"
    elif span2 is None:
        total, basis = span1, "Step 1 only (Step 2 ends before it starts)"
    else:
        total, basis = span1 + span2, "stamped per pass"
    if shared_by > 1:
        total, basis = round(total / shared_by, 1), f"{basis}; block average of {shared_by} rows (decision 2026-10-05)"
    return total, basis


def report(result):
    """Aggregates only: counts and issue labels, never a company name."""
    s = result.get("summary")
    print(f'FILE  {result["file"]}  sha256 {result["sha256"][:12]}…')
    print("      " + ", ".join(f"{k} {v}" for k, v in result["file_checks"].items()))
    if not s:
        print("      unreadable: stopped at the file checks")
        return
    print(f'ROWS  {s["rows"]} rows, {s["complete"]} Complete, {s["sdr_alone_rows"]} SDR alone, {s["with_agent_rows"]} With agent; '
          f'{s["prospects"]} Prospect, {s["dead_ends"]} Dead end; Pass 1 snapshot rows {s["snapshot_rows"]}')
    print(f'      pass-1 minutes readable on {s["pass1_minutes_readable"]} rows, pass-2 on {s["pass2_minutes_readable"]}; '
          f'rows with issues {s["rows_with_issues"]}')
    for label, n in sorted(result["issue_counts"].items(), key=lambda kv: (-kv[1], kv[0])):
        print(f"      {n} row(s): {label}")


def file_return(data, received, filed, received_on, draw, note=None, corrects=None):
    """Checks the two copies are identical, reads the filed copy, records the filing. A correction
    to a filed return is filed as the next version (-2, -3) with a note; a filed file is never edited."""
    folder = data / "sdr" / "in"
    filed_path = folder / filed
    assert filed_path.exists(), f"{filed} must already be in {folder}"
    log = data / "master-log" / "returns.jsonl"
    filed_before = [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line.strip()] if log.exists() else []
    if any(entry["sha256"] == sha256(filed_path) for entry in filed_before):
        raise SystemExit(f"{filed} is already filed (same hash). Nothing written.")
    if corrects:
        earlier = [entry for entry in filed_before if entry["filed_as"] == corrects]
        assert earlier, f"{corrects} is not a filed return, so nothing to correct"
        assert note, "a correction needs a note: what changed, who changed it, when"
        received, received_on = earlier[-1]["received_as"], earlier[-1]["received_on"]
        answers = earlier[-1]["answers"]
    else:
        received_path = folder / received
        assert received_path.exists(), f"{received} must be in {folder}, kept as received"
        assert sha256(received_path) == sha256(filed_path), "the filed copy is not byte-identical to the received file"
        # the builder's own list: strict release names, date then sequence, never a .FAILED copy
        from build_sdr_release import exchange_files
        releases = [p.name for p in exchange_files(data / "sdr" / "out", "out")]
        originals = [entry for entry in filed_before if not entry.get("corrects")]
        answers = releases[len(originals)] if len(originals) < len(releases) else None
    result = read_return(filed_path, draw=draw, release_path=(data / "sdr" / "out" / answers) if answers else None)
    report(result)
    record = {"kind": "SDR return", "received_as": received, "filed_as": filed, "sha256": result["sha256"],
              "received_on": received_on, "filed_at": datetime.now().astimezone().isoformat(timespec="seconds"),
              "answers": answers, "from": "SDR", "corrects": corrects, "note": note,
              "summary": result.get("summary"), "file_checks": result["file_checks"], "issue_counts": dict(result["issue_counts"])}
    with open(log, "a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    what = f"corrects {corrects}" if corrects else f"received as {received!r} on {received_on}"
    print(f"FILED {filed} ({what}), answers {answers}: recorded in {log.name}")
    return record


def made_up_return(path, companies, plant=False):
    """A small return on the contract's sheets. Clean by default: stamps in order, one snapshot
    per SDR-alone row, listed values. With plant=True, seven faults are planted: one per row for
    six rows, plus one Step 1 span stamped on two rows."""
    wb = Workbook()
    wb.active.title = SHEET_INSTRUCTIONS
    wb[SHEET_INSTRUCTIONS]["A1"] = "instructions"
    ws, snap, lists = wb.create_sheet(SHEET_RESEARCH), wb.create_sheet(SHEET_SNAPSHOT), wb.create_sheet(SHEET_LISTS)
    lists["A1"] = "lists"
    for sheet in (ws, snap):
        for c, name in enumerate(RESEARCH_COLUMNS, start=1):
            sheet.cell(row=HEADER_ROW, column=c, value=name)
    col = {name: i + 1 for i, name in enumerate(RESEARCH_COLUMNS)}
    for i, c in enumerate(companies):
        r = FIRST_DATA_ROW + i
        values = dict.fromkeys(RESEARCH_COLUMNS)
        values.update({"Order": c["work_order"], "Company ID": c["company_id"], "Company name": c["company_name"],
                       "Market": c["market"], "Arm": c["arm"], "Status": "Complete", "Date started": datetime(2026, 10, 1),
                       "Step 1 start": time(9, 0 + i), "Step 1 end": time(9, 30 + i), "Step 2 start": time(10, 0 + i), "Step 2 end": time(10, 20 + i),
                       "Verdict": "Prospect", "CRM status": "New", "Priority": "B"})
        for name in RESEARCH_FIELDS:
            values[name] = values.get(name) or ("Y" if name in PRODUCT_CATEGORIES else f"typed {name}")
        if plant:
            if i == 0:
                values["Step 1 end"] = time(8, 0)                       # end before start
            if i == 1:
                values["Step 2 start"], values["Step 2 end"] = values["Step 1 start"], values["Step 1 end"]   # copied stamps
            if i == 2:
                values["LGD"] = "Not Sure"                               # not Y or N
            if i == 3:
                values["Verdict"], values["Dead-end type"], values["Dead-end reason with source"] = "Dead end", None, "reason"  # type missing
                values["Priority"] = None
            if i == 4:
                values["Market"] = None                                  # grey cell blanked
            if i in (4, 5):                                              # one Step 1 span stamped on two rows
                values["Step 1 start"], values["Step 1 end"] = time(11, 0), time(12, 0)
        for name, value in values.items():
            ws.cell(row=r, column=col[name], value=value)
        if not (plant and i == 5):                                       # i == 5: snapshot missing
            snapshot = dict(values)
            snapshot["Step 2 start"] = snapshot["Step 2 end"] = None
            snapshot["Verdict"] = snapshot["CRM status"] = None
            snapshot["Contact names"] = None                             # pass 2 adds the contacts
            for name, value in snapshot.items():
                snap.cell(row=r, column=col[name], value=value)
    wb.save(path)


def self_test():
    results = []

    def check(name, passed, detail=""):
        results.append(bool(passed))
        print(f'{"PASS" if passed else "FAIL"}  {name}{"  [" + str(detail) + "]" if detail != "" else ""}')

    folder = Path(tempfile.mkdtemp(prefix="sdr-return-selftest-"))
    companies = [{"company_id": f"T{i:02d}", "company_name": f"Madeup {i} Jewels", "market": "USA" if i % 2 else "Gulf",
                  "arm": ARM_SDR, "work_order": i, "difficulty": "medium"} for i in range(1, 7)]
    draw = {"companies": companies}
    good, bad = folder / "good.xlsx", folder / "bad.xlsx"
    made_up_return(good, companies)
    made_up_return(bad, companies, plant=True)
    clean = read_return(good, draw=draw)
    check("positive control: a clean return has 0 issues on 6 rows", clean["summary"]["rows"] == 6 and clean["summary"]["rows_with_issues"] == 0,
          dict(clean["issue_counts"]))
    check("clean return: pass-1 and pass-2 minutes read on every row (30 and 20)",
          all(r["pass1_minutes"] == 30 and r["pass2_minutes"] == 20 for r in clean["rows"]))
    check("clean return: a snapshot per row, and pass 2 changed exactly one research field",
          all(r["snapshot"] and r["fields_changed_in_pass_2"] == 1 for r in clean["rows"]))
    planted = read_return(bad, draw=draw)
    want = {"Step 1: end before start", "Step 2: stamps equal Step 1 (one block, two passes)",
            "LGD: unsure value, counted as not found (decision 2026-10-05)", "Dead-end type: missing or not a listed value",
            "Market: grey cell changed (blank)", "Pass 1 snapshot: missing", "Step 1: stamps shared with 1 other row(s)"}
    found = set(planted["issue_counts"])
    check("negative control: all seven planted faults are found", want <= found, sorted(want - found))
    check("planted faults: 6 rows with issues, the copied stamps leave pass-2 minutes unreadable on that row",
          planted["summary"]["rows_with_issues"] == 6 and planted["summary"]["pass2_minutes_readable"] == 5)
    check("a row with end before start has no pass-1 minutes", planted["rows"][0]["pass1_minutes"] is None)
    check("clean return: SDR minutes per row are 50, read as stamped per pass",
          all(r["sdr_minutes"] == 50 and r["minutes_basis"] == "stamped per pass" for r in clean["rows"]))
    check("planted: Step 2 copied from Step 1 gives one span of 30 minutes for both passes",
          planted["rows"][1]["sdr_minutes"] == 30 and planted["rows"][1]["minutes_basis"].startswith("one span"))
    check("planted: two rows sharing one Step 1 span get the block average (80 minutes over 2 rows = 40 each)",
          all(planted["rows"][i]["sdr_minutes"] == 40 and "block average of 2 rows" in planted["rows"][i]["minutes_basis"] for i in (4, 5)))
    check("planted: the row that ends before it starts has no minutes and says why",
          planted["rows"][0]["sdr_minutes"] is None and planted["rows"][0]["minutes_basis"].startswith("unreadable"))
    # filing
    for sub in ("sdr/in", "sdr/out", "master-log"):
        (folder / sub).mkdir(parents=True)
    (folder / "sdr" / "out" / "sdr-2099-01-01-out.xlsx").write_bytes(good.read_bytes())
    (folder / "sdr" / "in" / "As Received.xlsx").write_bytes(good.read_bytes())
    (folder / "sdr" / "in" / "sdr-2099-01-02-in.xlsx").write_bytes(good.read_bytes())
    record = file_return(folder, "As Received.xlsx", "sdr-2099-01-02-in.xlsx", "2099-01-02", draw)
    check("filing: recorded against the release it answers, with the summary", record["answers"] == "sdr-2099-01-01-out.xlsx" and record["summary"]["rows"] == 6)
    try:
        file_return(folder, "As Received.xlsx", "sdr-2099-01-02-in.xlsx", "2099-01-02", draw)
        twice = False
    except SystemExit:
        twice = True
    check("filing: the same hash is refused a second time", twice)
    (folder / "sdr" / "in" / "sdr-2099-01-03-in.xlsx").write_bytes(bad.read_bytes())
    try:
        file_return(folder, "As Received.xlsx", "sdr-2099-01-03-in.xlsx", "2099-01-03", draw)
        mismatch = False
    except AssertionError:
        mismatch = True
    check("filing: a copy that is not byte-identical to the received file is refused", mismatch)
    (folder / "sdr" / "in" / "sdr-2099-01-02-in-2.xlsx").write_bytes(bad.read_bytes())
    record = file_return(folder, None, "sdr-2099-01-02-in-2.xlsx", None, draw, note="made-up correction", corrects="sdr-2099-01-02-in.xlsx")
    check("filing a correction: the next version, with its note, inheriting the release and the received name of the version it corrects",
          record["corrects"] == "sdr-2099-01-02-in.xlsx" and record["note"] == "made-up correction"
          and record["answers"] == "sdr-2099-01-01-out.xlsx" and record["received_as"] == "As Received.xlsx")
    print(f"\n{sum(results)} of {len(results)} self-test checks passed")
    return all(results)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if "--self-test" in sys.argv:
        sys.exit(0 if self_test() else 1)
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("mode", choices=["file", "check"])
    parser.add_argument("path", nargs="?", help="check: the workbook to read")
    parser.add_argument("--received", help="file: the name the SDR gave the file, as kept in sdr/in/")
    parser.add_argument("--filed", help="file: the contract name, sdr-YYYY-MM-DD-in.xlsx, already copied into sdr/in/")
    parser.add_argument("--received-on", help="file: the date it was received, YYYY-MM-DD")
    parser.add_argument("--correction-of", help="file: this is the next version of an already filed return (give its filed name); "
                                                "the identity check is skipped and --note is required")
    parser.add_argument("--note", help="file: what changed, who changed it, when (required for a correction)")
    parser.add_argument("--draw", help='the draw record; default: the one draw-*.json in master-log/; "none" for the slice file')
    parser.add_argument("--data-root", default=str(HERE / "data-exports"))
    args = parser.parse_args()
    data = Path(args.data_root)
    if args.draw == "none":          # a file outside the draw, such as the slice file
        draw_path = None
    else:
        draw_path = Path(args.draw) if args.draw else next(iter(sorted((data / "master-log").glob("draw-*.json"))), None)
    draw = json.loads(draw_path.read_text(encoding="utf-8")) if draw_path else None
    if args.mode == "check":
        report(read_return(args.path, draw=draw))
        return
    if args.correction_of:
        assert args.filed and args.note, "a correction needs --filed, --correction-of and --note"
    else:
        assert args.received and args.filed and args.received_on, "file needs --received, --filed and --received-on"
    file_return(data, args.received, args.filed, args.received_on, draw, note=args.note, corrects=args.correction_of)


if __name__ == "__main__":
    main()
