"""Build the master log workbook from its sources. Mayank and the runner only: never sent.

    python build_master_log.py --draw data-exports/master-log/draw-2026-10-01.json \
        --agent-runs data-exports/agent-runs/batch-v1

The workbook is a view. Its sources are the draw record, the release log (releases.jsonl), the
filed SDR returns (returns.jsonl, written by sdr_return.py), the agent runs and the lead's filed
returns (lead-returns.jsonl, read with lead_return.py). Every build writes a new file,
data-exports/master-log/master-log-YYYY-MM-DD.xlsx (then -2, -3 on the same day): no version is
overwritten. Sheets: Header, Companies, Rows, Reserves, Balance, Releases, Returns, Lead returns.
The SDR's Step 1 and Step 2 stamps are filed as pass 1 and pass 2 on SDR-alone rows and as check
and completing on Pair rows; a later return overrides an earlier one for the rows it holds. Nothing
the SDR typed is changed; the reader's issues sit in "Ingest notes".
Agent runs: an Agent alone row takes its company's run in --agent-runs (pass only the v1 batch: one
v1 run per company serves that arm); a Pair row takes the agent run it was built on, as the release
log names it (v1 or v2), with the row file's hash checked. The lead's scores reach their rows
through the lead key; a row takes them from the latest lead return in which it carries all three.
"""

import argparse
import hashlib
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

import lead_return
import sdr_return
from contract import (
    ARM_AGENT, ARM_SDR, DIFFICULTIES, EARLY, FIELD_KEYS, LATE, MARKETS, MASTER_LOG_CONTRACT, RESEARCH_FIELDS,
    ROW_AGENT_ALONE, ROW_PAIR, ROW_SDR_ALONE, VERDICT_DEAD_END,
)
from xlsx_util import BASE, BOLD, TITLE, WARN_FILL, WRAP_TOP, gate, next_version_path, write_table

HERE = Path(__file__).resolve().parent
INGEST_EXTRA = ["SDR status", "Date started", "Pass-1 minutes", "Pass-2 minutes", "Check minutes", "Completing minutes",
                "SDR minutes (both passes)", "Minutes basis", "Pass-1 snapshot", "Fields changed in pass 2", "Ingest notes"]
AGENT_EXTRA = ["Agent row file", "Tool calls", "Web searches", "Web fetches"]
LEAD_EXTRA = ["Lead file", "Lead row code", "Lead scoring start", "Lead scoring end", "Lead comment",
              "Dead end right (lead re-check)", "Lead re-check comment", "Lead re-check file"]
ROWS_HEADERS = (
    ["Row ID", "Row type"] + MASTER_LOG_CONTRACT + RESEARCH_FIELDS
    + ["Dead-end reason with source", "CRM status", "Source file"] + INGEST_EXTRA + AGENT_EXTRA + LEAD_EXTRA
)
STAMP_HEADERS = ["Pass-1 start", "Pass-1 end", "Pass-2 start", "Pass-2 end", "Check start", "Check end",
                 "Completing start", "Completing end", "Management asked at", "Management answer added at",
                 "Lead scoring start", "Lead scoring end"]
AGENT_TIMES = ["Agent start", "Agent end"]
AGENT_RUN_FIELDS = ["Prompt version", "Model", "Effort", "Agent start", "Agent end", "Agent minutes", "Apollo credits",
                    "Flag value", "Validator result", "Retry count", "Agent row file", "Tool calls"]
LEAD_SCORE_FIELDS = ["Acceptable", "Correctness", "Completeness", "Lead scoring minutes", "Lead file", "Lead row code"]
ROW_SUFFIX = {ROW_AGENT_ALONE: "A", ROW_SDR_ALONE: "S", ROW_PAIR: "P"}


def planned_rows(companies):
    """Two rows per company: the agent-alone row, then the SDR-alone or the pair row."""
    rows = []
    for c in sorted(companies, key=lambda c: c["work_order"]):
        for row_type in (ROW_AGENT_ALONE, ROW_SDR_ALONE if c["arm"] == ARM_SDR else ROW_PAIR):
            row = dict.fromkeys(ROWS_HEADERS)
            row.update({
                "Row ID": f'{c["company_id"]}-{ROW_SUFFIX[row_type]}', "Row type": row_type,
                "Company ID": c["company_id"], "Company name": c["company_name"], "Market": c["market"],
                "Difficulty": c["difficulty"], "Arm": c["arm"], "Early or late": c["early_or_late"],
                "Work order": c["work_order"],
            })
            rows.append(row)
    return rows


def load_returns(data, draw):
    """The filed SDR returns, in filing order, each read with sdr_return.read_return."""
    log = data / "master-log" / "returns.jsonl"
    records = [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line.strip()] if log.exists() else []
    results = [sdr_return.read_return(data / "sdr" / "in" / rec["filed_as"], draw=draw) for rec in records]
    return records, results


def apply_returns(rows_by_id, results):
    """Files every company row of every return onto its master-log row; a later return wins."""
    filed = 0
    for result in results:
        for r in result["rows"]:
            row_type = ROW_SDR_ALONE if r["arm"] == ARM_SDR else ROW_PAIR
            target = rows_by_id.get(f'{r["company_id"]}-{ROW_SUFFIX[row_type]}')
            if target is None:
                continue
            s = r["stamps"]
            pass1 = None if r.get("pass1_minutes_shared") else r["pass1_minutes"]
            if row_type == ROW_SDR_ALONE:
                stamps = {"Pass-1 start": s["Step 1 start"], "Pass-1 end": s["Step 1 end"], "Pass-2 start": s["Step 2 start"],
                          "Pass-2 end": s["Step 2 end"], "Pass-1 minutes": pass1, "Pass-2 minutes": r["pass2_minutes"]}
            else:
                stamps = {"Check start": s["Step 1 start"], "Check end": s["Step 1 end"], "Completing start": s["Step 2 start"],
                          "Completing end": s["Step 2 end"], "Check minutes": pass1, "Completing minutes": r["pass2_minutes"]}
            target.update(stamps)
            target.update({name: r["fields"][name] for name in RESEARCH_FIELDS})
            started = r["date_started"]
            target.update({
                "Management asked at": s["Management asked at"], "Management answer added at": s["Management answer added at"],
                "Verdict": r["verdict"], "Dead-end category": r["dead_end_type"], "Dead-end reason with source": r["dead_end_reason"],
                "CRM status": r["crm_status"], "Notes": r["notes"], "Source file": result["file"],
                "SDR status": r["status"], "Date started": started.date() if isinstance(started, datetime) else started,
                "SDR minutes (both passes)": r.get("sdr_minutes"), "Minutes basis": r.get("minutes_basis"),
                "Pass-1 snapshot": ("Yes" if r["snapshot"] else "No") if "snapshot" in r else None,
                "Fields changed in pass 2": r.get("fields_changed_in_pass_2"),
                "Ingest notes": "; ".join(r["issues"]) or None,
            })
            filed += 1
    return filed


def local_time(stamp):
    """An ISO stamp from a run record as a local datetime without its zone (Excel keeps none), or None."""
    return datetime.fromisoformat(stamp).replace(tzinfo=None) if stamp else None


def agent_run(folder):
    """(validated agent row, run record) in one company folder; (None, None) when it holds no row."""
    folder = Path(folder)
    if not (folder / "row.json").exists():
        return None, None
    run_path = folder / "run.json"
    return (json.loads((folder / "row.json").read_text(encoding="utf-8")),
            json.loads(run_path.read_text(encoding="utf-8")) if run_path.exists() else {})


def agent_fields(row, run, row_file):
    """What an agent row files on its master-log row: its run (version, model, effort, times, credits, validator,
    retries, tool calls) and its flag."""
    tools = (run.get("summary") or {}).get("tools_called") or {}
    flag = row["flag"]
    return {
        "Prompt version": run.get("prompt_version"), "Model": run.get("model_asked"), "Effort": run.get("effort"),
        "Agent start": local_time(run.get("agent_start")), "Agent end": local_time(run.get("agent_end")),
        "Agent minutes": run.get("agent_minutes"),
        # the web-only versions call no Apollo tool: 0 credits. A run that did is left blank, to be read from its credit blocks
        "Apollo credits": (0 if not run.get("apollo_calls_that_ran") else None) if run else None,
        "Flag value": "Yes" if flag["likely_dead_end"] == "yes" else "No",
        "Flag reason": flag["reason"].strip() or None, "Flag source": flag["source"].strip() or None,
        "Validator result": run.get("validator_result"), "Retry count": run.get("retry_count"), "Agent row file": row_file,
        "Tool calls": len(run["tool_calls"]) if "tool_calls" in run else None,
        "Web searches": tools.get("WebSearch", 0) if run else None, "Web fetches": tools.get("WebFetch", 0) if run else None,
    }


def apply_agent_runs(rows_by_id, companies, data, alone_folder, releases):
    """Agent alone rows take the run in alone_folder: its values, verdict and flag. Pair rows take the run of the agent
    row they were built on, as the release log names it, when that row file's hash is unchanged; the Pair row's values
    and verdict stay the SDR's. Returns counts."""
    counts = Counter()
    if alone_folder:
        for c in companies:
            row, run = agent_run(Path(alone_folder) / c["company_id"])
            target = rows_by_id.get(f'{c["company_id"]}-{ROW_SUFFIX[ROW_AGENT_ALONE]}')
            if row is None or target is None:
                continue
            row_file = f'{Path(alone_folder).name}/{c["company_id"]}/row.json'
            verdict = row["verdict"]
            dead = verdict["value"] == VERDICT_DEAD_END
            target.update(agent_fields(row, run, row_file))
            target.update({name: row["fields"][key]["value"].strip() or None for name, key in FIELD_KEYS.items()})
            target.update({
                "Verdict": verdict["value"], "Dead-end category": verdict["dead_end_type"] if dead else None,
                "Dead-end reason with source": " ".join(p for p in (verdict["reason"].strip(), verdict["source"].strip()) if p) if dead else None,
                "CRM status": row.get("crm_status"), "Source file": row_file,
            })
            counts["agent alone"] += 1
    for entry in releases:
        for record in entry.get("agent_rows") or []:
            target = rows_by_id.get(f'{record["company_id"]}-{ROW_SUFFIX[ROW_PAIR]}')
            path = Path(data) / "agent-runs" / record["row"]
            if target is None:
                continue
            if not path.exists() or hashlib.sha256(path.read_bytes()).hexdigest() != record["row_sha256"]:
                counts["pair agent rows missing or changed since their release"] += 1
                continue
            row, run = agent_run(path.parent)
            target.update(agent_fields(row, run, record["row"]))
            counts["pair"] += 1
    return counts


def load_lead_returns(data):
    """The filed lead returns, in filing order, each read with lead_return.read_return."""
    log = Path(data) / "master-log" / "lead-returns.jsonl"
    records = [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line.strip()] if log.exists() else []
    return records, [lead_return.read_return(Path(data) / "lead" / "in" / rec["filed_as"], data, rec["answers"]) for rec in records]


def apply_lead_returns(rows_by_id, results):
    """The lead's scores onto their master-log rows, through the key. A row takes them from the latest return in which
    it carries all three scores; an unscored return changes nothing."""
    filed = 0
    for result in results:
        for r in result["rows"]:
            target = rows_by_id.get(f'{r["company_id"]}-{ROW_SUFFIX.get(r["row_type"], "?")}')
            if not r["scored"] or target is None:
                continue
            target.update({"Acceptable": r["acceptable"], "Correctness": r["correctness"], "Completeness": r["completeness"],
                           "Lead scoring minutes": r["minutes"], "Lead file": result["file"], "Lead row code": r["row_code"],
                           "Lead scoring start": r["start"], "Lead scoring end": r["end"], "Lead comment": r["comment"]})
            filed += 1
    return filed


def apply_lead_rechecks(rows_by_id, results):
    """The lead's dead-end re-check answers (DECISIONS.md 18) onto their rows, beside the first scores, which stay as typed."""
    filed = 0
    for result in results:
        for r in result["rows"]:
            target = rows_by_id.get(f'{r["company_id"]}-{ROW_SUFFIX.get(r["row_type"], "?")}')
            if not r.get("recheck") or target is None:
                continue
            target.update({"Dead end right (lead re-check)": r["recheck"], "Lead re-check comment": r["comment"],
                           "Lead re-check file": result["file"]})
            filed += 1
    return filed


def count_table(companies, row_attr, row_labels, col_attr, col_labels):
    body = []
    for label in row_labels:
        counts = [sum(1 for c in companies if c[row_attr] == label and c[col_attr] == col) for col in col_labels]
        if sum(counts):
            body.append([label] + counts + [sum(counts)])
    totals = [sum(1 for c in companies if c[col_attr] == col) for col in col_labels]
    return body + [["Total"] + totals + [sum(totals)]]


def build(draw, draw_name, releases, out, returns=((), ()), data=None, alone_folder=None, lead=((), ())):
    records, results = returns
    lead_records, lead_results = lead
    companies = sorted(draw["companies"], key=lambda c: c["work_order"])
    rows = planned_rows(companies)
    rows_by_id = {row["Row ID"]: row for row in rows}
    apply_returns(rows_by_id, results)
    agent_counts = apply_agent_runs(rows_by_id, companies, data, alone_folder, releases)
    lead_filed = apply_lead_returns(rows_by_id, lead_results)
    rechecked = apply_lead_rechecks(rows_by_id, lead_results)
    wb = Workbook()

    ws = wb.active
    ws.title = "Header"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    ws["A1"] = "MASTER LOG: SDR research agent, measured trial (capstone)"
    ws["A1"].font = TITLE
    ws["A2"] = "Mayank and the runner only. Never send this file: it holds the arms, the difficulty labels and every arm's rows."
    ws["A2"].font, ws["A2"].fill = BOLD, WARN_FILL
    ws.merge_cells("A2:B2")
    mix = Counter(c["difficulty"] for c in companies)
    lines = [
        ("This file", out.name),
        ("Built at", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Draw record", draw_name),
        ("Seed", draw["seed"]),
        ("Drawn at", draw["drawn_at"]),
        ("List file (SHA-256)", f'{Path(draw["list_file"]).name} ({draw["list_sha256"]})'),
        ("Draw rule", draw["rule"]),
    ]
    lines += [(f"Recorded with the draw ({i})", note) for i, note in enumerate(draw["notes"], start=1)]
    lines += [
        ("Where the difficulty label lives", "In this file only. It is never on the SDR's sheet or the lead's sheet."),
        ("Mix of the 30 drawn", ", ".join(f"{d} {mix[d]}" for d in DIFFICULTIES)
         + f'. Pool of {draw["pool_size"]}: ' + ", ".join(
             f"{d} {sum(1 for c in draw['companies'] + draw['not_drawn'] if c['difficulty'] == d)}" for d in DIFFICULTIES)),
        ("Arms", f"{sum(c['arm'] == ARM_SDR for c in companies)} SDR alone ({sum(c['early_or_late'] == EARLY for c in companies)} early), "
                 f"{sum(c['arm'] == ARM_AGENT for c in companies)} With agent"),
        ("Rows sheet", "One line per scored row: an Agent alone row for every company, plus its SDR alone or Pair row. "
                       "Step 1 and Step 2 stamps from the SDR are filed as pass 1 and pass 2 on SDR alone rows, "
                       "and as check and completing on Pair rows."),
        ("SDR returns filed", f"{len(records)} (see the Returns sheet). The SDR's values are filed as typed; the reader's "
                              "issues per row sit in the Rows sheet's Ingest notes. Minutes are left blank where the stamps "
                              "cannot give them (shared block stamps, end before start, Step 2 equal to Step 1)."),
        ("Agent runs filed", f'{agent_counts["agent alone"]} Agent alone rows from {Path(alone_folder).name if alone_folder else "no folder"}; '
                             f'{agent_counts["pair"]} Pair rows from the agent row each was built on, as the release log names it '
                             f'(missing or changed since their release: {agent_counts["pair agent rows missing or changed since their release"]}).'),
        ("Lead returns filed", f"{len(lead_records)} (see the Lead returns sheet); {lead_filed} scored rows filed through the lead key. "
                               "Lead scoring minutes are whole minutes between the lead's two stamps (0 = within the minute). "
                               f"Dead-end re-check answers filed: {rechecked} (DECISIONS.md 18)."),
    ]
    for r, (key, value) in enumerate(lines, start=4):
        ws.cell(row=r, column=1, value=key).font = BOLD
        cell = ws.cell(row=r, column=2, value=value)
        cell.font, cell.alignment = BASE, WRAP_TOP
        ws.cell(row=r, column=1).alignment = WRAP_TOP

    ws = wb.create_sheet("Companies")
    write_table(
        ws, ["Work order", "Company ID", "Company name", "Market", "Difficulty", "CRM (as pasted)", "Arm", "Early or late", "List no"],
        [[c["work_order"], c["company_id"], c["company_name"], c["market"], c["difficulty"], c["crm"], c["arm"],
          c["early_or_late"], c["list_no"]] for c in companies],
        widths=[8, 10, 44, 9, 11, 12, 12, 11, 8],
    )
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Rows")
    widths = ([10, 12] + [14] * len(MASTER_LOG_CONTRACT) + [18] * len(RESEARCH_FIELDS) + [30, 14, 30]
              + [12] * (len(INGEST_EXTRA) - 1) + [60] + [24, 9, 9, 9] + [24, 9, 10, 10, 60, 12, 50, 24])
    widths[ROWS_HEADERS.index("Minutes basis")] = 44
    widths[ROWS_HEADERS.index("Company name")] = 40
    write_table(ws, ROWS_HEADERS, [[row[h] for h in ROWS_HEADERS] for row in rows], widths=widths)
    for r in range(2, len(rows) + 2):
        for name in STAMP_HEADERS:
            ws.cell(row=r, column=ROWS_HEADERS.index(name) + 1).number_format = "h:mm"
        for name in AGENT_TIMES:
            ws.cell(row=r, column=ROWS_HEADERS.index(name) + 1).number_format = "yyyy-mm-dd h:mm:ss"
        ws.cell(row=r, column=ROWS_HEADERS.index("Date started") + 1).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "E2"

    ws = wb.create_sheet("Reserves")
    queue = []
    for market in MARKETS:
        position = 0
        for r in sorted((r for r in draw["not_drawn"] if r["market"] == market), key=lambda r: r["reserve_front_no"]):
            position += 1
            queue.append([market, position, r["company_name"], "Not drawn from the list", r["difficulty"], r["crm"]])
        for r in sorted((r for r in draw["pasted_reserves"] if r["market"] == market), key=lambda r: r["reserve_no"]):
            position += 1
            queue.append([market, position, r["company_name"], "Pasted reserve", None, None])
    last = write_table(ws, ["Market", "Queue position", "Company name", "Source", "Difficulty", "CRM (as pasted)"], queue,
                       widths=[9, 10, 44, 24, 11, 12])
    note = ws.cell(row=last + 2, column=1, value="Use only if a name falls, in queue order within the market. "
                                                 "Pasted reserves carry no difficulty label yet: the lead labels one before it is used.")
    note.font = BASE

    ws = wb.create_sheet("Balance")
    ws.column_dimensions["A"].width = 16
    row = 1
    sdr = [c for c in companies if c["arm"] == ARM_SDR]
    everyone = draw["companies"] + draw["not_drawn"]
    for c in everyone:
        c["_drawn"] = "Drawn" if c in draw["companies"] else "Not drawn"
        c["_cell"] = f'{c["market"]} {c["difficulty"]}'
    cells = [f"{m} {d}" for m in MARKETS for d in DIFFICULTIES]
    for title, items, row_attr, row_labels, col_attr, col_labels in [
        ("Arms by market", companies, "market", MARKETS, "arm", [ARM_SDR, ARM_AGENT]),
        ("Arms by difficulty", companies, "difficulty", DIFFICULTIES, "arm", [ARM_SDR, ARM_AGENT]),
        ("SDR-alone companies, early or late, by market", sdr, "market", MARKETS, "early_or_late", [EARLY, LATE]),
        ("SDR-alone companies, early or late, by difficulty", sdr, "difficulty", DIFFICULTIES, "early_or_late", [EARLY, LATE]),
        ("Pool and draw, by market and difficulty", everyone, "_cell", cells, "_drawn", ["Drawn", "Not drawn"]),
    ]:
        ws.cell(row=row, column=1, value=title).font = BOLD
        row = write_table(ws, [""] + col_labels + ["Total"], count_table(items, row_attr, row_labels, col_attr, col_labels),
                          widths=None, start_row=row + 1) + 2
    for col in "BCD":
        ws.column_dimensions[col].width = 12

    ws = wb.create_sheet("Releases")
    keys = ["file", "built_at", "to", "built_on", "through_order", "rows", "new_rows", "sdr_alone_rows", "with_agent_rows",
            "agent_rows_on_sdr_alone_companies", "sha256"]
    write_table(
        ws, ["File", "Built at", "To", "Built on", "Through order", "Rows", "New rows", "SDR-alone rows", "With-agent rows",
             "Agent rows on SDR-alone companies (must be 0)", "SHA-256"],
        [[entry.get(k) for k in keys] for entry in releases], widths=[28, 24, 8, 26, 9, 7, 7, 9, 9, 18, 66],
    )
    ws = wb.create_sheet("Returns")
    write_table(
        ws, ["Filed as", "Received as", "Received on", "Filed at", "Answers", "Corrects", "Note", "Rows", "Complete", "Prospect",
             "Dead end", "Pass 1 snapshot rows", "Rows with issues", "Pass-1 minutes readable", "Pass-2 minutes readable",
             "Rows with SDR minutes", "SDR minutes total", "Issue counts", "SHA-256"],
        [[rec["filed_as"], rec["received_as"], rec["received_on"], rec["filed_at"], rec["answers"], rec.get("corrects"), rec.get("note"),
          *((rec.get("summary") or {}).get(k) for k in ("rows", "complete", "prospects", "dead_ends", "snapshot_rows", "rows_with_issues",
                                                        "pass1_minutes_readable", "pass2_minutes_readable", "sdr_minutes_rows",
                                                        "sdr_minutes_total")),
          "; ".join(f"{n} {label}" for label, n in sorted(rec["issue_counts"].items(), key=lambda kv: (-kv[1], kv[0]))),
          rec["sha256"]] for rec in records],
        widths=[24, 30, 12, 26, 24, 24, 50, 6, 9, 9, 9, 10, 10, 12, 12, 12, 12, 80, 66],
    )
    ws = wb.create_sheet("Lead returns")
    write_table(
        ws, ["Filed as", "Received as", "Received on", "Filed at", "Answers", "Checked against", "Rows", "Scored rows",
             "Scored rows by row type", "Re-check answers", "Checks failed", "Note", "SHA-256"],
        [[rec["filed_as"], rec["received_as"], rec["received_on"], rec["filed_at"], rec["answers"], result.get("checked_against"),
          result["summary"]["rows"], result["summary"]["scored_rows"],
          ", ".join(f"{k} {v}" for k, v in sorted(result["summary"]["scored_rows_by_type"].items())) or None,
          ", ".join(f"{k} {v}" for k, v in sorted((result["summary"].get("answers") or {}).items())) or None,
          ", ".join(name for name, passed in result["checks"].items() if not passed) or "none", rec.get("note"), rec["sha256"]]
         for rec, result in zip(lead_records, lead_results)],
        widths=[24, 40, 12, 26, 28, 26, 7, 8, 34, 16, 24, 80, 66],
    )
    wb.save(out)


def verify(out, draw, releases, returns=((), ()), data=None, alone_folder=None, lead=((), ())):
    records, filed = returns
    lead_records, lead_results = lead
    wb = load_workbook(out)
    results = []

    def check(name, passed, detail=""):
        results.append(bool(passed))
        print(f'{"PASS" if passed else "FAIL"}  {name}{"  [" + str(detail) + "]" if detail != "" else ""}')

    check("sheets", wb.sheetnames == ["Header", "Companies", "Rows", "Reserves", "Balance", "Releases", "Returns", "Lead returns"],
          wb.sheetnames)
    header_text = " ".join(str(c.value) for row in wb["Header"].iter_rows() for c in row if c.value is not None)
    check("header records every note from the draw, the seed and the list hash",
          all(note in header_text for note in draw["notes"]) and str(draw["seed"]) in header_text and draw["list_sha256"] in header_text,
          f'{len(draw["notes"])} notes')
    ws = wb["Companies"]
    got = [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2)]
    want = [(c["work_order"], c["company_id"], c["company_name"], c["market"], c["difficulty"], c["crm"], c["arm"],
             c["early_or_late"], c["list_no"]) for c in sorted(draw["companies"], key=lambda c: c["work_order"])]
    check("Companies sheet equals the draw record, row for row", got == want, f"{len(got)} rows")
    ws = wb["Rows"]
    headers = [c.value for c in ws[1]]
    check("Rows header holds every field of the master-log contract",
          all(name in headers for name in MASTER_LOG_CONTRACT), f"{len(headers)} columns")
    kinds = Counter(ws.cell(row=r, column=headers.index("Row type") + 1).value for r in range(2, ws.max_row + 1))
    per_company = Counter(ws.cell(row=r, column=headers.index("Company ID") + 1).value for r in range(2, ws.max_row + 1))
    check("Rows: 30 agent-alone, 15 SDR-alone, 15 pair, two per company",
          kinds == {ROW_AGENT_ALONE: 30, ROW_SDR_ALONE: 15, ROW_PAIR: 15} and set(per_company.values()) == {2}, dict(kinds))
    check("Reserves: not-drawn first, then the pasted reserves",
          wb["Reserves"].max_row - 3 == len(draw["not_drawn"]) + len(draw["pasted_reserves"]),
          f'{len(draw["not_drawn"])} + {len(draw["pasted_reserves"])}')
    check("Releases sheet equals the release log", wb["Releases"].max_row - 1 == len(releases), f"{len(releases)} releases")
    check("Returns sheet equals the returns log", wb["Returns"].max_row - 1 == len(records), f"{len(records)} returns")
    ws = wb["Rows"]
    by_id = {ws.cell(row=r, column=headers.index("Row ID") + 1).value: r for r in range(2, ws.max_row + 1)}
    col = {name: headers.index(name) + 1 for name in headers}
    expected = {}
    for result in filed:
        for r in result["rows"]:
            expected[f'{r["company_id"]}-{"S" if r["arm"] == ARM_SDR else "P"}'] = (r, result["file"])
    mismatches = []
    for row_id, (r, source) in expected.items():
        line = by_id.get(row_id)
        if line is None:
            mismatches.append((row_id, "no master-log row"))
            continue
        first, last = ("Pass-1 start", "Pass-2 end") if r["arm"] == ARM_SDR else ("Check start", "Completing end")
        got = (ws.cell(row=line, column=col["Verdict"]).value, ws.cell(row=line, column=col["CRM status"]).value,
               ws.cell(row=line, column=col[first]).value, ws.cell(row=line, column=col[last]).value,
               ws.cell(row=line, column=col["Source file"]).value, *(ws.cell(row=line, column=col[name]).value for name in RESEARCH_FIELDS))
        want = (r["verdict"], r["crm_status"], r["stamps"]["Step 1 start"], r["stamps"]["Step 2 end"], source,
                *(r["fields"][name] for name in RESEARCH_FIELDS))
        if got != want:
            mismatches.append((row_id, "values differ"))
    check("every row of every filed return is on its master-log row: verdict, CRM, first and last stamp, source file, 25 fields",
          not mismatches and (len(expected) > 0 or not filed), f"{len(expected)} rows, mismatches {mismatches[:3]}")
    stamp_cells = [ws.cell(row=line, column=col[name]) for line in by_id.values() for name in STAMP_HEADERS
                   if ws.cell(row=line, column=col[name]).value is not None]
    check("filed stamps are time cells shown as h:mm", all(c.number_format == "h:mm" and hasattr(c.value, "hour") for c in stamp_cells),
          f"{len(stamp_cells)} stamp cells")
    touched = [line for line in by_id.values() if ws.cell(row=line, column=col["Source file"]).value is None
               and any(ws.cell(row=line, column=col[name]).value is not None for name in ("Verdict", "Pass-1 start", "Check start"))]
    check("no row carries values without a source file", not touched, len(touched))

    def value(row_id, name):
        return ws.cell(row=by_id[row_id], column=col[name]).value

    if alone_folder:
        with_runs = [c["company_id"] for c in draw["companies"] if (Path(alone_folder) / c["company_id"] / "row.json").exists()]
        complete = [cid for cid in with_runs if all(value(f"{cid}-A", name) is not None for name in AGENT_RUN_FIELDS)
                    and value(f"{cid}-A", "Verdict") is not None]
        check("every Agent alone row with a validated run carries it: version, model, effort, start, end, minutes, credits, flag, "
              "validator, retries, tool calls, verdict", len(complete) == len(with_runs), f"{len(complete)} of {len(with_runs)}")
    given = {rec["company_id"]: rec for entry in releases for rec in entry.get("agent_rows") or []}
    carried = [cid for cid, rec in given.items() if value(f"{cid}-P", "Agent row file") == rec["row"]
               and value(f"{cid}-P", "Prompt version") == rec["prompt_version"]]
    check("every Pair row carries the agent run the release log says it was built on, its row file unchanged (v1 or v2)",
          len(carried) == len(given), f"{len(carried)} of {len(given)}: " + ", ".join(
              f"{v} {n}" for v, n in sorted(Counter(rec["prompt_version"] for rec in given.values()).items())))
    want = {}
    for result in lead_results:
        for r in result["rows"]:
            if r["scored"]:
                want[f'{r["company_id"]}-{ROW_SUFFIX[r["row_type"]]}'] = (r["acceptable"], r["correctness"], r["completeness"],
                                                                          r["minutes"], result["file"], r["row_code"])
    got = {row_id: tuple(value(row_id, name) for name in LEAD_SCORE_FIELDS) for row_id in by_id
           if value(row_id, "Acceptable") is not None}
    check("the lead's scores on the master log equal the filed lead returns, row for row through the key", got == want,
          f"{len(want)} scored rows; " + ", ".join(f"{k} {v}" for k, v in sorted(Counter(
              ws.cell(row=by_id[row_id], column=col["Row type"]).value for row_id in want).items())))
    want_re = {f'{r["company_id"]}-{ROW_SUFFIX[r["row_type"]]}': (r["recheck"], result["file"])
               for result in lead_results for r in result["rows"] if r.get("recheck")}
    got_re = {row_id: (value(row_id, "Dead end right (lead re-check)"), value(row_id, "Lead re-check file")) for row_id in by_id
              if value(row_id, "Dead end right (lead re-check)") is not None}
    check("the dead-end re-check answers on the master log equal the filed re-check returns, row for row", got_re == want_re,
          f"{len(want_re)} answers")
    check("Lead returns sheet equals the lead-returns log", wb["Lead returns"].max_row - 1 == len(lead_records),
          f"{len(lead_records)} lead returns")
    fonts = {cell.font.name for s in wb.worksheets for row in s.iter_rows() for cell in row if cell.value is not None}
    formulas = sum(1 for s in wb.worksheets for row in s.iter_rows() for cell in row
                   if isinstance(cell.value, str) and cell.value.startswith("="))
    check("Arial only, no formulas", fonts == {"Arial"} and formulas == 0, f"{sorted(fonts)}, {formulas} formulas")
    print(f"\n{sum(results)} of {len(results)} checks passed")
    return all(results)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--draw", required=True)
    parser.add_argument("--data-root", default=str(HERE / "data-exports"), help="for tests on made-up data")
    parser.add_argument("--agent-runs", help="the run folder whose runs serve the agent-alone arm (the v1 batch)")
    args = parser.parse_args()
    data = Path(args.data_root)
    log_dir = data / "master-log"
    draw = json.loads(Path(args.draw).read_text(encoding="utf-8"))
    log = log_dir / "releases.jsonl"
    releases = [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line.strip()] if log.exists() else []
    returns = load_returns(data, draw)
    lead = load_lead_returns(data)
    failed = [(rec["filed_as"], name) for rec, result in zip(*lead) for name, passed in result["checks"].items() if not passed]
    if failed:
        raise SystemExit(f"a filed lead return fails its read checks, so its scores are not filed: {failed}")
    out = next_version_path(log_dir, f"master-log-{datetime.now():%Y-%m-%d}")
    build(draw, Path(args.draw).name, releases, out, returns, data, args.agent_runs, lead)
    print(f"MASTER LOG {out.name}: {len(returns[0])} SDR return(s), {len(lead[0])} lead return(s) filed onto the Rows sheet; "
          f"agent runs from {Path(args.agent_runs).name if args.agent_runs else 'no folder'} and the release log\n")
    draw = json.loads(Path(args.draw).read_text(encoding="utf-8"))
    ok = verify(out, draw, releases, returns, data, args.agent_runs, lead)
    ok = gate(out) and ok
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
