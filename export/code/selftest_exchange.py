"""End-to-end self-test of the offline exchange on MADE-UP data, in a temporary folder.

    python selftest_exchange.py [--keep]

Nothing real is read or written. A made-up list is drawn; a first SDR release is built; a made-up
SDR return is filed (re-saved through Excel, as a real return would be); made-up agent rows are
written; a second release adds the With-agent rows on top of the return; the lead's blind scoring
file is built, scored, returned and extended. Along the way the guards are attacked:
  - a release while the last one has not come back,
  - a With-agent company with no agent row,
  - an agent flag planted on an SDR-alone row (the contamination count),
  - a lead view that would not mix all three arms,
  - a planted marker source, a planted flag reason and a planted CRM value, which must each
    appear 0 times in the lead's file (spec exit condition 7).
With --keep the folder is left in place for a look.
"""

import csv
import json
import os
import shutil
import stat
import subprocess
import sys
import tempfile
from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

import build_lead_view as lead
import build_sdr_release as release
from contract import ARM_AGENT, ARM_SDR, FIELD_KEYS, PRODUCT_CATEGORIES, RESEARCH_COLUMNS, RESEARCH_FIELDS, SHEET_RESEARCH, SHEET_SNAPSHOT
from validate_row import example_dead_end, example_row, validate

HERE = Path(__file__).resolve().parent
RESULTS = []
PLANTED_SOURCE = "https://planted-marker-source.example/Q7"
PLANTED_FLAG = "PLANTED-FLAG-REASON-Q7"
PLANTED_NOTE = "PLANTED-NOTE-Q7"
PLANTED_CRM = "Dormant"


def check(name, passed, detail=""):
    RESULTS.append(bool(passed))
    print(f'{"PASS" if passed else "FAIL"}  {name}{"  [" + str(detail) + "]" if detail != "" else ""}')


def run(script, *args):
    done = subprocess.run([sys.executable, str(HERE / script), *map(str, args)], capture_output=True, text=True, encoding="utf-8")
    return done.returncode, done.stdout + done.stderr


def last_line(text):
    lines = text.strip().splitlines()
    return lines[-1][:90] if lines else ""


def made_up_list(folder):
    mix = [("USA", "hard", 10), ("USA", "medium", 12), ("USA", "easy", 2), ("Gulf", "hard", 9), ("Gulf", "medium", 5), ("Gulf", "easy", 2)]
    rows, n = [], 0
    for market, difficulty, count in mix:
        for i in range(count):
            n += 1
            rows.append([n, f"Madeup {market} {difficulty} {i + 1:02d} Jewels", market, difficulty, "New"])
    with open(folder / "companies.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["list_no", "company_name", "market", "difficulty", "crm"])
        writer.writerows(rows)
    with open(folder / "reserves.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["market", "reserve_no", "company_name"])
        writer.writerows([["USA", 1, "Madeup reserve U1 Jewels"], ["Gulf", 1, "Madeup reserve G1 Jewels"]])


def excel_resave(source, target):
    """Open in Excel and save a copy, so the return carries Excel's own XML like a real one."""
    script = (
        "$e = New-Object -ComObject Excel.Application; $e.Visible = $false; $e.DisplayAlerts = $false; "
        f"$w = $e.Workbooks.Open('{source}'); $w.SaveAs('{target}', 51); $w.Close($false); $e.Quit(); "
        "[void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($e)"
    )
    subprocess.run(["powershell.exe", "-NoProfile", "-Command", script], capture_output=True, text=True, timeout=180)
    return Path(target).exists()


def made_up_sdr_return(out_file, in_file, leave_open=0):
    """What the SDR would send back. SDR-alone rows are researched in full. With-agent rows keep
    the agent's cells and get a verdict; the first one becomes a confirmed dead end. One row gets
    the planted CRM value. The last `leave_open` rows get a verdict but are not marked Complete.
    Every row is stamped (Step 1 of 30 minutes, Step 2 of 20) and every Complete SDR-alone row is
    copied to the Pass 1 snapshot sheet as values, the way the SDR is asked to."""
    wb = load_workbook(out_file)
    ws, snap = wb[SHEET_RESEARCH], wb[SHEET_SNAPSHOT]
    columns = release.header_columns(ws)
    rows = release.data_rows(ws, columns)
    first_pair = True
    for i, (r, _) in enumerate(rows):
        def put(name, value):
            ws.cell(row=r, column=columns[name], value=value)

        arm = ws.cell(row=r, column=columns["Arm"]).value
        open_row = i >= len(rows) - leave_open
        put("Date started", datetime(2099, 1, 1))
        put("Step 1 start", time(9, i))
        put("Step 1 end", time(9, 30 + i))
        if not open_row:
            put("Step 2 start", time(10, i))
            put("Step 2 end", time(10, 20 + i))
        if arm == ARM_SDR:
            typed = {"Parent company": "Madeup Holdings", "Detailed research": "- Typed by the SDR\n- Sells matched pairs",
                     "Gold": "Y", "Natural diamond": "Y", "Main country": "USA", "Priority": "B"}
            for name in RESEARCH_FIELDS:
                put(name, typed.get(name, "N" if name in PRODUCT_CATEGORIES else f"typed {name}"))
        put("Verdict", "Prospect")
        if arm == ARM_AGENT and first_pair:
            first_pair = False
            put("Verdict", "Dead end")
            put("Dead-end type", "Competitor")
            put("Dead-end reason with source", "Owned by a competitor group https://example.com/group")
        put("CRM status", PLANTED_CRM if i == 0 else "New")
        put("Notes", PLANTED_NOTE)
        put("Status", "Step 1 done" if open_row else "Complete")
        if arm == ARM_SDR and not open_row:
            for name in RESEARCH_COLUMNS:
                value = ws.cell(row=r, column=columns[name]).value
                if name in ("Step 2 start", "Step 2 end", "Verdict", "CRM status", "Contact names"):
                    value = None
                snap.cell(row=r, column=columns[name], value=value)
    staged = Path(in_file).with_suffix(".staged.xlsx")
    wb.save(staged)
    resaved = excel_resave(staged, in_file)
    staged.unlink()
    return resaved


def made_up_agent_rows(companies, folder):
    """A validated agent row for every company. The first carries the planted source and flag
    reason; every second Gulf company is a lab-grown-only dead end."""
    flag_planted = False
    for i, c in enumerate(companies):
        row = example_dead_end() if c["market"] == "Gulf" and i % 2 == 0 else example_row()
        row.update({"company_id": c["company_id"], "company_name": c["company_name"], "market": c["market"]})
        if i == 0:
            row["fields"]["parent_company"]["source"] = PLANTED_SOURCE
        if not flag_planted and row["flag"]["likely_dead_end"] == "no":
            row["flag"]["reason"], flag_planted = PLANTED_FLAG, True
        assert not validate(row), validate(row)
        (folder / c["company_id"]).mkdir(parents=True)
        (folder / c["company_id"] / "row.json").write_text(json.dumps(row), encoding="utf-8")
        run = {"company_id": c["company_id"], "prompt_version": "v1", "model_asked": "made-up-model", "effort": "high",
               "agent_start": "2099-01-01T10:00:00+05:30", "agent_end": "2099-01-01T10:02:00+05:30", "agent_minutes": 2.0,
               "validator_result": "valid", "retry_count": 0, "apollo_calls_that_ran": {},
               "summary": {"tools_called": {"WebSearch": 2, "WebFetch": 3}}, "tool_calls": [{"tool": "WebSearch"}] * 5}
        (folder / c["company_id"] / "run.json").write_text(json.dumps(run), encoding="utf-8")


def whole_file_text(path):
    wb = load_workbook(path)
    cells = [cell for ws in wb.worksheets for row in ws.iter_rows() for cell in row]
    return " \n ".join([str(c.value) for c in cells if c.value is not None] + [c.comment.text for c in cells if c.comment])


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    root = Path(tempfile.mkdtemp(prefix="sdr-exchange-selftest-"))
    for sub in ("list", "master-log", "sdr/out", "sdr/in", "lead/out", "lead/in", "agent-runs/batch"):
        (root / sub).mkdir(parents=True)
    made_up_list(root / "list")
    draw_path = root / "master-log" / "draw.json"
    code, _ = run("draw_arms.py", "--list", root / "list" / "companies.csv", "--reserves", root / "list" / "reserves.csv",
                  "--seed", 777, "--write", "--out", draw_path)
    check("made-up draw written", code == 0)
    draw = json.loads(draw_path.read_text(encoding="utf-8"))
    ordered = sorted(draw["companies"], key=lambda c: c["work_order"])
    common = ["--draw", draw_path, "--data-root", root]
    runs = root / "agent-runs/batch"

    print("\n--- SDR exchange")
    code, text = run("build_sdr_release.py", *common, "--through-order", 7, "--date", "2099-01-01")
    check("release 1 (7 SDR-alone rows, on the template) passes its checks and the gate", code == 0, last_line(text))
    code, text = run("build_sdr_release.py", *common, "--through-order", 9, "--date", "2099-01-02")
    check("guard: a second release is refused while the first has not come back", code != 0 and "has not come back" in text)
    check("made-up SDR return 1 filed (re-saved through Excel)",
          made_up_sdr_return(root / "sdr/out/sdr-2099-01-01-out.xlsx", root / "sdr/in/sdr-2099-01-01-in.xlsx"))

    print("\n--- Ingest into the master log")
    (root / "sdr/in/As received 1.xlsx").write_bytes((root / "sdr/in/sdr-2099-01-01-in.xlsx").read_bytes())
    code, text = run("sdr_return.py", "file", "--received", "As received 1.xlsx", "--filed", "sdr-2099-01-01-in.xlsx",
                     "--received-on", "2099-01-01", "--draw", draw_path, "--data-root", root)
    check("return 1 recorded in the returns log with its checks (0 issues on clean made-up rows)",
          code == 0 and "FILED" in text and "rows with issues 0" in text, last_line(text))
    code, text = run("build_master_log.py", "--draw", draw_path, "--data-root", root)
    check("master log built on the draw, the release log and the filed return: all checks and the gate", code == 0, last_line(text))
    if code:
        print(text)
    ws = load_workbook(sorted((root / "master-log").glob("master-log-*.xlsx"))[-1])["Rows"]
    column = {c.value: i + 1 for i, c in enumerate(ws[1])}
    filed = [r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=column["Source file"]).value == "sdr-2099-01-01-in.xlsx"]
    check("the 7 SDR-alone rows carry pass-1 and pass-2 stamps, 30 and 20 minutes, and a snapshot that pass 2 added one field to",
          len(filed) == 7 and all(ws.cell(row=r, column=column["Pass-1 minutes"]).value == 30
                                  and ws.cell(row=r, column=column["Pass-2 minutes"]).value == 20
                                  and ws.cell(row=r, column=column["Pass-1 snapshot"]).value == "Yes"
                                  and ws.cell(row=r, column=column["Fields changed in pass 2"]).value == 1
                                  and ws.cell(row=r, column=column["Row type"]).value == "SDR alone" for r in filed), len(filed))
    check("no other master-log row was touched", sum(1 for r in range(2, ws.max_row + 1)
                                                      if ws.cell(row=r, column=column["Verdict"]).value is not None) == 7)

    wb = load_workbook(root / "sdr/in/sdr-2099-01-01-in.xlsx")
    ws = wb[SHEET_RESEARCH]
    ws.cell(row=3, column=release.header_columns(ws)["Market"]).value = None     # a grey cell the SDR cleared
    pasted = ws.cell(row=4, column=release.header_columns(ws)["Emails"])
    pasted.value = "pasted@example.com"                                          # a value the SDR pasted in Calibri
    pasted.font = Font(name="Calibri", size=11)
    wb.save(root / "sdr/in/sdr-2099-01-01-in.xlsx")

    through = 14
    code, text = run("build_sdr_release.py", *common, "--through-order", through, "--date", "2099-01-02")
    check("guard: a With-agent company with no agent row is refused", code != 0 and "no validated agent row" in text)
    made_up_agent_rows(ordered, runs)
    code, text = run("build_sdr_release.py", *common, "--through-order", through, "--date", "2099-01-02", "--agent-runs", runs)
    check("release 2 (With-agent rows added on top of the return) passes its checks and the gate", code == 0, last_line(text))
    if code:
        print(text)

    second = root / "sdr/out/sdr-2099-01-02-out.xlsx"
    logged = [json.loads(line) for line in (root / "master-log/releases.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
    check("release 2 restored the grey Market cell the return had blanked, and logged it",
          logged[-1].get("grey_cells_restored") == 1 and logged[-2].get("grey_cells_restored") == 0, logged[-1].get("grey_cells_restored"))
    check("release 2 kept the SDR's pasted Calibri cell as returned, value and font, and logged it",
          logged[-1].get("sdr_cells_kept_in_another_font") == 1, logged[-1].get("sdr_cells_kept_in_another_font"))
    for label, path in (("release 1", root / "sdr/out/sdr-2099-01-01-out.xlsx"), ("release 2", second)):
        lines = [c.value for c in load_workbook(path)[release.SHEET_INSTRUCTIONS]["A"]]
        check(f"{label} prints the rules added on 5 Oct once on its Instructions sheet",
              lines.count(release.RULES_ADDED[1]) == 1 and all(line in lines for line in release.RULES_ADDED[1:]))
    pairs = [c for c in ordered if c["work_order"] <= through and c["arm"] == ARM_AGENT]
    wb = load_workbook(second)
    ws = wb[SHEET_RESEARCH]
    columns = release.header_columns(ws)
    rows = dict((cid, r) for r, cid in release.data_rows(ws, columns))
    check("the SDR's typed rows survive the next release untouched",
          all(ws.cell(row=rows[c["company_id"]], column=columns["Notes"]).value == PLANTED_NOTE for c in ordered[:7]))
    pair_cells = [ws.cell(row=rows[c["company_id"]], column=columns[name]) for c in pairs for name in FIELD_KEYS]
    check("With-agent rows: every research cell carries a marker colour",
          all(release.marker_colour(cell) in release.MARKER_COLOURS.values() for cell in pair_cells), f"{len(pair_cells)} cells")
    check("contamination count on the good file is 0", release.agent_content_on_sdr_alone(ws, columns) == 0)
    planted = ws.cell(row=rows[pairs[0]["company_id"]], column=columns["Emails"])
    planted.font = Font(name="Calibri", size=11)
    if planted.value is None:
        planted.value = "written by the release"
    ours, kept = release.font_outliers(wb, load_workbook(root / "sdr/in/sdr-2099-01-01-in.xlsx"))
    check("negative control: a Calibri cell in a row this release wrote is caught, the SDR's own one is not",
          len(ours) == 1 and len(kept) == 1, f"{len(ours)} caught, {len(kept)} kept")
    victim = next(c for c in ordered if c["arm"] == ARM_SDR)
    ws.cell(row=rows[victim["company_id"]], column=columns["Agent flag"], value="Yes")
    check("negative control: an agent flag planted on an SDR-alone row is counted",
          release.agent_content_on_sdr_alone(ws, columns) == 1)
    new_pairs = [c for c in pairs if c["work_order"] > 7]
    check("release 2 logs where each agent row it added came from: run folder and the row file's hash",
          logged[-1]["agent_runs"] == ["batch"] and [a["company_id"] for a in logged[-1]["agent_rows"]] == [c["company_id"] for c in new_pairs]
          and all(a["row_sha256"] == release.sha256(runs / a["company_id"] / "row.json") for a in logged[-1]["agent_rows"]),
          f'{len(logged[-1]["agent_rows"])} rows')
    corrected = root / "sdr/in/sdr-2099-01-01-in-2.xlsx"
    wb = load_workbook(root / "sdr/in/sdr-2099-01-01-in.xlsx")
    ws = wb[SHEET_RESEARCH]
    ws.cell(row=3, column=release.header_columns(ws)["Notes"]).value = PLANTED_NOTE + " (corrected)"
    wb.save(corrected)
    code, text = run("sdr_return.py", "file", "--filed", corrected.name, "--correction-of", "sdr-2099-01-01-in.xlsx",
                     "--note", "made-up correction", "--draw", draw_path, "--data-root", root)
    check("a correction of return 1 is filed as its next version, answering release 1",
          code == 0 and "answers sdr-2099-01-01-out.xlsx" in text, last_line(text))
    code, text = run("build_sdr_release.py", *common, "--through-order", through + 2, "--date", "2099-01-03", "--agent-runs", runs)
    check("guard: release 3 is refused while release 2 has not come back, though sdr/in/ now holds two files for two releases",
          code != 0 and "has not come back" in text and not (root / "sdr/out/sdr-2099-01-03-out.xlsx").exists(), last_line(text))

    print("\n--- Lead exchange")
    code, text = run("build_lead_view.py", *common, "--date", "2099-01-03")
    check("guard: no lead file while no company has both its rows (no agent rows given)", code != 0 and "no whole company" in text,
          last_line(text))
    check("made-up SDR return 2 filed, two rows not yet Complete",
          made_up_sdr_return(second, root / "sdr/in/sdr-2099-01-02-in.xlsx", leave_open=2))
    wb = load_workbook(root / "sdr/in/sdr-2099-01-02-in.xlsx")
    ws = wb[SHEET_RESEARCH]
    ws.cell(row=3, column=release.header_columns(ws)["Gold"]).value = "Not Sure"     # typed by the SDR in a Y/N cell
    unsure_company = ws.cell(row=3, column=release.header_columns(ws)["Company ID"]).value
    wb.save(root / "sdr/in/sdr-2099-01-02-in.xlsx")
    code, text = run("build_lead_view.py", *common, "--date", "2099-01-03", "--agent-runs", runs)
    check("lead release 1 passes its checks and the gate", code == 0, last_line(text))
    if code:
        print(text)
    first_lead = root / "lead/out/lead-2099-01-03-out.xlsx"
    key = json.loads((root / "master-log/lead-key.json").read_text(encoding="utf-8"))
    companies = {}
    for entry in key["rows"]:
        companies.setdefault(entry["company_id"], []).append(entry["row_type"])
    check("lead release 1 holds the 12 whole companies (24 rows); the 18 with only an agent row wait",
          len(key["rows"]) == 2 * (through - 2) and len(companies) == through - 2 and all(lead.is_whole(k) for k in companies.values()),
          f'{len(key["rows"])} rows, {len(companies)} companies')
    flagged = next(p.parent.name for p in runs.glob("*/row.json") if PLANTED_FLAG in p.read_text(encoding="utf-8"))
    sourced = next(p.parent.name for p in runs.glob("*/row.json") if PLANTED_SOURCE in p.read_text(encoding="utf-8"))
    check("the rows carrying the planted values are in the lead's file, so the 0 counts below test something",
          {flagged, sourced, ordered[0]["company_id"], unsure_company} <= set(companies))
    logged = [json.loads(line) for line in (root / "master-log/releases.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
    ws_lead = load_workbook(root / "lead/out/lead-2099-01-03-out.xlsx")[lead.SHEET]
    code_of = {e["company_id"]: e["row_code"] for e in key["rows"] if e["row_type"] == lead.ROW_SDR_ALONE}
    row_of = {code: r for r, code in lead.scoring_rows(ws_lead)}
    gold = ws_lead.cell(row=row_of[code_of[unsure_company]], column=lead.HEADERS.index("Gold") + 1).value
    check("the SDR's 'Not Sure' in a Y/N cell is shown blank, and counted in the release log",
          gold is None and logged[-1].get("unsure_cells_shown_blank") == 1, f'{gold!r}, {logged[-1].get("unsure_cells_shown_blank")}')
    inputs = "".join(p.read_text(encoding="utf-8") for p in runs.glob("*/row.json")) + whole_file_text(root / "sdr/in/sdr-2099-01-02-in.xlsx")
    check("positive control: all four planted values are present in what the lead view was built from, and the same scan finds them",
          all(token in inputs for token in (PLANTED_SOURCE, PLANTED_FLAG, PLANTED_CRM, PLANTED_NOTE)))
    text = whole_file_text(first_lead)
    for label, token in (("planted marker source", PLANTED_SOURCE), ("planted flag reason", PLANTED_FLAG),
                         ("planted CRM value", PLANTED_CRM), ("planted note", PLANTED_NOTE)):
        check(f"negative control: the {label} appears 0 times in the lead's file", text.count(token) == 0, text.count(token))
    ws = load_workbook(first_lead)[lead.SHEET]
    column = {name: i + 1 for i, name in enumerate(lead.HEADERS)}
    dead = [r for r, _ in lead.scoring_rows(ws) if ws.cell(row=r, column=column["Verdict"]).value == "Dead end"]
    extra = sum(ws.cell(row=r, column=column[name]).value is not None for r in dead for name in RESEARCH_FIELDS)
    check("dead-end rows from both arms show only the short-row fields", len(dead) >= 2 and extra == 0, f"{len(dead)} rows, {extra} extra cells")

    code, text = run("build_lead_view.py", *common, "--reissue", first_lead.name, "--reason", "made-up: the printed lines changed")
    reissued = root / "lead/out/lead-2099-01-03-out-2.xlsx"
    codes_of = lambda path: [c for _, c in lead.scoring_rows(load_workbook(path)[lead.SHEET])]
    entry = [json.loads(line) for line in (root / "master-log/releases.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()][-1]
    check("an unworked lead file is issued again: the same codes in the same order, the first file kept, the log marks it superseded",
          code == 0 and reissued.exists() and first_lead.exists() and codes_of(reissued) == codes_of(first_lead)
          and entry.get("reissue_of") == first_lead.name, last_line(text))
    wb = load_workbook(reissued)
    ws = wb[lead.SHEET]
    scored = lead.scoring_rows(ws)[:5]
    for r, _ in scored:
        for name, typed in (("Acceptable", "Yes"), ("Correctness", "None"), ("Completeness", "Full"), ("Comment", "typed by the lead")):
            ws.cell(row=r, column=column[name], value=typed)
        ws.cell(row=r, column=column["Scoring start"], value=time(10, 0))
        ws.cell(row=r, column=column["Scoring end"], value=time(10, 2))
    staged = root / "lead/in/staged.xlsx"
    wb.save(staged)
    check("made-up lead return filed (re-saved through Excel)", excel_resave(staged, root / "lead/in/lead-2099-01-03-in.xlsx"))
    staged.unlink()
    (root / "lead/in/As received lead 1.xlsx").write_bytes((root / "lead/in/lead-2099-01-03-in.xlsx").read_bytes())
    code, text = run("lead_return.py", "file", "--received", "As received lead 1.xlsx", "--filed", "lead-2099-01-03-in.xlsx",
                     "--received-on", "2099-01-03", "--data-root", root)
    check("the lead's return is filed, answering the re-issued release; every check passes; 5 rows scored",
          code == 0 and "answers lead-2099-01-03-out-2.xlsx" in text and "FAIL" not in text and "5 scored" in text, last_line(text))
    code, text = run("lead_return.py", "file", "--received", "As received lead 1.xlsx", "--filed", "lead-2099-01-03-in.xlsx",
                     "--received-on", "2099-01-03", "--data-root", root)
    check("guard: the same lead return is never filed twice", code != 0 and "already filed" in text, last_line(text))
    code, text = run("build_master_log.py", "--draw", draw_path, "--data-root", root, "--agent-runs", runs)
    check("master log with the agent runs and the lead's scores: all checks and the gate", code == 0, last_line(text))
    if code:
        print(text)
    # the newest version: by name, "-2.xlsx" sorts before ".xlsx"
    ws_log = load_workbook(max((root / "master-log").glob("master-log-*.xlsx"), key=lambda p: p.stat().st_mtime))["Rows"]
    log_col = {c.value: i + 1 for i, c in enumerate(ws_log[1])}
    log_rows = {ws_log.cell(row=r, column=log_col["Row ID"]).value: r for r in range(2, ws_log.max_row + 1)}
    by_code = {e["row_code"]: e for e in json.loads((root / "master-log/lead-key.json").read_text(encoding="utf-8"))["rows"]}
    suffix = {lead.ROW_AGENT_ALONE: "A", lead.ROW_SDR_ALONE: "S", lead.ROW_PAIR: "P"}
    targets = [f'{by_code[c]["company_id"]}-{suffix[by_code[c]["row_type"]]}' for _, c in scored]
    check("the 5 scores sit on the master-log rows the key names, with 2 lead minutes each, and on no other row",
          all(ws_log.cell(row=log_rows[t], column=log_col["Acceptable"]).value == "Yes"
              and ws_log.cell(row=log_rows[t], column=log_col["Lead scoring minutes"]).value == 2 for t in targets)
          and sum(ws_log.cell(row=r, column=log_col["Acceptable"]).value is not None for r in log_rows.values()) == 5)
    with_agent = [c["company_id"] for c in ordered if c["arm"] == ARM_AGENT and c["work_order"] <= through]
    check("every Agent alone row carries its run; each Pair row released so far carries the agent run it was built on",
          all(ws_log.cell(row=log_rows[f'{c["company_id"]}-A'], column=log_col["Agent minutes"]).value == 2.0 for c in ordered)
          and all(ws_log.cell(row=log_rows[f"{cid}-P"], column=log_col["Agent row file"]).value == f"batch/{cid}/row.json"
                  for cid in with_agent if cid not in {c["company_id"] for c in ordered[:7]}))
    code, text = run("build_lead_view.py", *common, "--date", "2099-01-04", "--agent-runs", runs)
    check("guard: no new whole company, no release (the lead's return changes nothing)", code != 0 and "no whole company" in text,
          last_line(text))
    made_up_sdr_return(second, root / "sdr/in/sdr-2099-01-03-in.xlsx")
    code, text = run("build_lead_view.py", *common, "--date", "2099-01-04", "--agent-runs", runs)
    check("lead release 2 (2 more whole companies, on the lead's return) passes its checks and the gate", code == 0, last_line(text))
    if code:
        print(text)
    ws = load_workbook(root / "lead/out/lead-2099-01-04-out.xlsx")[lead.SHEET]
    kept = [ws.cell(row=r, column=column["Comment"]).value for r, _ in scored]
    check("the lead's typed scores survive the next release", kept == ["typed by the lead"] * 5)
    check("lead release 2 holds every whole company once: 14 companies, 28 rows", len(lead.scoring_rows(ws)) == 2 * through,
          len(lead.scoring_rows(ws)))
    code, text = run("build_lead_view.py", *common, "--reissue", "lead-2099-01-04-out.xlsx", "--reason", "made-up")
    check("guard: a lead file carrying scores is never issued again", code != 0 and "carries scores" in text, last_line(text))
    sent = sorted((root / "sdr/out").glob("sdr-*-out.xlsx")) + sorted((root / "lead/out").glob("lead-*-out*.xlsx"))
    check("every sent file is sealed read-only once built", sent and not any(os.access(f, os.W_OK) for f in sent), f"{len(sent)} files")

    print("\n--- Dead-end re-check (DECISIONS.md 18)")
    second_lead = root / "lead/out/lead-2099-01-04-out.xlsx"
    wb = load_workbook(second_lead)
    ws = wb[lead.SHEET]
    for r, _ in lead.scoring_rows(ws):
        for name, typed in (("Acceptable", "No"), ("Correctness", "None"), ("Completeness", "Full")):
            if not ws.cell(row=r, column=column[name]).value:
                ws.cell(row=r, column=column[name], value=typed)
        ws.cell(row=r, column=column["Scoring start"], value=time(11, 0))
        ws.cell(row=r, column=column["Scoring end"], value=time(11, 1))
    staged = root / "lead/in/staged.xlsx"
    wb.save(staged)
    excel_resave(staged, root / "lead/in/lead-2099-01-04-in.xlsx")
    staged.unlink()
    (root / "lead/in/As received lead 2.xlsx").write_bytes((root / "lead/in/lead-2099-01-04-in.xlsx").read_bytes())
    code, text = run("lead_return.py", "file", "--received", "As received lead 2.xlsx", "--filed", "lead-2099-01-04-in.xlsx",
                     "--received-on", "2099-01-04", "--data-root", root)
    check("lead release 2 comes back scored and is filed, every check passing", code == 0 and "FAIL" not in text, last_line(text))
    ws = load_workbook(root / "lead/in/lead-2099-01-04-in.xlsx")[lead.SHEET]
    dead = [code for r, code in lead.scoring_rows(ws) if ws.cell(row=r, column=column["Verdict"]).value == "Dead end"]
    code, text = run("build_lead_view.py", *common, "--date", "2099-01-05", "--recheck-dead-ends")
    recheck = root / "lead/out/lead-2099-01-05-out.xlsx"
    check("the dead-end re-check passes its checks and the gate, and is sealed read-only",
          code == 0 and recheck.exists() and not os.access(recheck, os.W_OK), last_line(text))
    wb = load_workbook(recheck)
    ws = wb[lead.RECHECK_SHEET]
    first = lead.RECHECK_HEADER_ROW + 1
    codes = [ws.cell(row=r, column=1).value for r in range(first, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    check("the re-check holds every dead-end row of the return, by its code, in order, and nothing else", codes == dead and dead,
          f"{len(codes)} rows")
    question = lead.RECHECK_HEADERS.index(lead.RECHECK_QUESTION) + 1
    for i, r in enumerate(range(first, first + len(codes))):
        ws.cell(row=r, column=question, value="Yes" if i == 0 else "No")
        ws.cell(row=r, column=question + 1, value=None if i == 0 else "made-up: sells natural diamonds too")
    wb.save(staged)
    excel_resave(staged, root / "lead/in/lead-2099-01-05-in.xlsx")
    staged.unlink()
    (root / "lead/in/As received recheck.xlsx").write_bytes((root / "lead/in/lead-2099-01-05-in.xlsx").read_bytes())
    code, text = run("lead_return.py", "file", "--received", "As received recheck.xlsx", "--filed", "lead-2099-01-05-in.xlsx",
                     "--received-on", "2099-01-05", "--data-root", root)
    check("the re-check comes back and is filed as the return of the re-check release",
          code == 0 and "answers lead-2099-01-05-out.xlsx" in text and "FAIL" not in text, last_line(text))
    code, text = run("build_master_log.py", "--draw", draw_path, "--data-root", root, "--agent-runs", runs)
    check("master log with the re-check: all checks and the gate", code == 0, last_line(text))
    if code:
        print(text)
    ws_log = load_workbook(max((root / "master-log").glob("master-log-*.xlsx"), key=lambda p: p.stat().st_mtime))["Rows"]
    log_col = {c.value: i + 1 for i, c in enumerate(ws_log[1])}
    log_rows = {ws_log.cell(row=r, column=log_col["Row ID"]).value: r for r in range(2, ws_log.max_row + 1)}
    by_code = {e["row_code"]: e for e in json.loads((root / "master-log/lead-key.json").read_text(encoding="utf-8"))["rows"]}
    want = {f'{by_code[c]["company_id"]}-{suffix[by_code[c]["row_type"]]}': ("Yes" if i == 0 else "No") for i, c in enumerate(codes)}
    got = {rid: ws_log.cell(row=r, column=log_col["Dead end right (lead re-check)"]).value for rid, r in log_rows.items()
           if ws_log.cell(row=r, column=log_col["Dead end right (lead re-check)"]).value}
    check("the re-check answers sit on the rows the key names, beside the first scores, and on no other row", got == want,
          f"{len(got)} answers")

    print(f"\n{sum(RESULTS)} of {len(RESULTS)} self-test checks passed. Folder: {root}")
    if "--keep" not in sys.argv:
        for folder, _, files in os.walk(root):          # the sealed made-up releases: clear the bit so the folder can go
            for name in files:
                os.chmod(Path(folder) / name, stat.S_IWRITE | stat.S_IREAD)
        shutil.rmtree(root, ignore_errors=True)
    sys.exit(0 if all(RESULTS) else 1)


if __name__ == "__main__":
    main()
